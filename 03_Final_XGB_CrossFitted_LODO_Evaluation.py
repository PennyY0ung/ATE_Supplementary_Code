from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
import time
import warnings
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd

from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.feature_selection import SelectPercentile, f_classif
from sklearn.impute import SimpleImputer
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    cohen_kappa_score,
    confusion_matrix,
    f1_score,
    precision_recall_fscore_support,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder
from sklearn.utils.class_weight import compute_sample_weight

try:
    from xgboost import XGBClassifier
except Exception as exc:
    raise ImportError(
        "13-0固定使用XGB，请先安装xgboost，例如：pip install xgboost"
    ) from exc

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception as exc:
    raise ImportError("需要openpyxl用于Excel输出，请先安装openpyxl") from exc


# =============================================================================
# 1. 固定路径与实验定义
# =============================================================================

SCRIPT_VERSION = "13-0_v2_final_20260821"

DEFAULT_BASE_DIR = Path(__file__).resolve().parent

INPUT_NAME = "Analytical_Dataset_1001.xlsx"

DIR_12B_NAME = (
    "12B_TargetedBayesian_Stability_NestedLODO_FiveModels"
)
FILE_12B_NAME = (
    "12B_TargetedBayesian_Stability_NestedLODO_FiveModels.xlsx"
)

RUN_DIR_NAME = "13-0_XGB_CrossFittedLODO_Evaluation_V2"
OUTPUT_XLSX_NAME = "13-0_XGB_CrossFittedLODO_Evaluation_V2.xlsx"
CONFIG_JSON_NAME = "13-0_XGB_V2_Run_Configuration.json"
LOG_NAME = "13-0_XGB_V2_Run.log"

SHEET_MAP = {"Summer": "summer", "Winter": "winter"}
EXPECTED_N = {"Summer": 477, "Winter": 524}
EXPECTED_GROUPS = {"Summer": 190, "Winter": 296}

TARGET = "Thermal_satisfaction_3"
RAW_TO_MODEL = {1: 0, 2: 1, 3: 2}
MODEL_TO_RAW = {0: 1, 1: 2, 2: 3}
CLASS_ORDER = [0, 1, 2]

OBJECTIVE = [
    "Temperature",
    "Wind_speed",
    "Humidity",
    "Globe_temperature",
    "Illumination",
]
SPATIAL = ["Location", "Scene"]
SUBJECTIVE = [
    "TSV", "TCV", "TPV", "TAV",
    "SSV", "SCV", "SPV", "SAV",
    "HSV", "WSV", "HPV", "WPV",
]
EP_NUMERIC = OBJECTIVE + SUBJECTIVE
EP_CATEGORICAL = SPATIAL
EP_FEATURES = EP_NUMERIC + EP_CATEGORICAL

COMMON_FAMILY = "XGB"

# 必须与12B冻结值一致
CANDIDATE_MODELS_12B = ["GB", "CAT", "ET", "LGBM", "XGB"]
RANDOM_STATE = 20260816
MODEL_N_JOBS = -1
OUTER_FIT_PHASE_OFFSET = 900000


# =============================================================================
# 2. 日志、进度与辅助
# =============================================================================

class Tee:
    def __init__(self, *streams):
        self.streams = streams

    def write(self, data):
        for s in self.streams:
            s.write(data)
            s.flush()

    def flush(self):
        for s in self.streams:
            s.flush()


def fmt_seconds(seconds: float) -> str:
    if not np.isfinite(seconds):
        return "--"
    seconds = max(float(seconds), 0.0)
    if seconds < 60:
        return f"{seconds:.0f}s"
    if seconds < 3600:
        m, s = divmod(int(seconds), 60)
        return f"{m}m{s:02d}s"
    h, rem = divmod(int(seconds), 3600)
    m, s = divmod(rem, 60)
    return f"{h}h{m:02d}m{s:02d}s"


def ascii_bar(done: int, total: int, width: int = 30) -> str:
    if total <= 0:
        return "[" + "." * width + "]"
    frac = min(max(done / total, 0), 1)
    n = int(round(frac * width))
    if done >= total:
        return "[" + "=" * width + "]"
    n = min(n, width - 1)
    return "[" + "=" * n + ">" + "." * (width - n - 1) + "]"


def print_progress(done: int, total: int, start_time: float):
    elapsed = time.time() - start_time
    if done > 0:
        eta = elapsed / done * max(total - done, 0)
        finish = datetime.now() + timedelta(seconds=eta)
        finish_txt = finish.strftime("%Y-%m-%d %H:%M:%S")
    else:
        eta = float("nan")
        finish_txt = "--"

    print(
        f"[总体进度] {ascii_bar(done, total)} "
        f"{done}/{total} folds ({done/total*100 if total else 0:.1f}%) "
        f"| 已耗时 {fmt_seconds(elapsed)} "
        f"| ETA {fmt_seconds(eta)} "
        f"| 预计完成 {finish_txt}"
    )


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def bool_from_excel(v: Any) -> bool:
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    s = str(v).strip().lower()
    return s in {"true", "1", "yes", "y"}


def resolve_xlsx(exact: Path, folder: Path, keyword: str) -> Path:
    if exact.exists():
        return exact

    if not folder.exists():
        raise FileNotFoundError(f"找不到目录：{folder}")

    candidates = [
        p for p in folder.glob("*.xlsx")
        if keyword.lower() in p.name.lower() and not p.name.startswith("~$")
    ]
    if len(candidates) == 1:
        return candidates[0]

    all_xlsx = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]
    if len(all_xlsx) == 1:
        return all_xlsx[0]

    raise FileNotFoundError(
        f"无法唯一确定{keyword}结果Excel。\n"
        f"目录：{folder}\n"
        f"候选：{[p.name for p in all_xlsx]}"
    )


# =============================================================================
# 3. 数据读取与硬检查 —— 与12B一致
# =============================================================================

def load_data(input_file: Path) -> dict[str, pd.DataFrame]:
    raw = pd.read_excel(
        input_file,
        sheet_name=list(SHEET_MAP.values()),
        engine="openpyxl",
    )

    required = ["Date", "measurement_group", TARGET] + EP_FEATURES
    out: dict[str, pd.DataFrame] = {}

    for season, sheet in SHEET_MAP.items():
        df = raw[sheet].copy()
        df.columns = [str(c).strip() for c in df.columns]

        missing = [c for c in required if c not in df.columns]
        if missing:
            raise KeyError(f"{season}缺少字段：{missing}")

        if len(df) != EXPECTED_N[season]:
            raise ValueError(
                f"{season}: N={len(df)}，预期={EXPECTED_N[season]}"
            )

        df["Date"] = pd.to_datetime(
            df["Date"], errors="coerce"
        ).dt.normalize()

        df[TARGET] = pd.to_numeric(df[TARGET], errors="coerce")

        for c in EP_NUMERIC:
            df[c] = pd.to_numeric(df[c], errors="coerce")

        for c in EP_CATEGORICAL:
            df[c] = df[c].astype(str).str.strip()

        df["measurement_group"] = (
            df["measurement_group"].astype(str).str.strip()
        )

        if df[required].isna().any().any():
            bad = df[required].isna().sum()
            raise ValueError(
                f"{season}存在缺失/非法值："
                f"{bad[bad > 0].to_dict()}"
            )

        df[TARGET] = df[TARGET].astype(int)

        if set(df[TARGET].unique()) != {1, 2, 3}:
            raise ValueError(
                f"{season}: Thermal_satisfaction_3必须完整为1/2/3"
            )

        if df["Date"].nunique() != 4:
            raise ValueError(
                f"{season}: 调查日={df['Date'].nunique()}，应为4天"
            )

        if df["measurement_group"].nunique() != EXPECTED_GROUPS[season]:
            raise ValueError(
                f"{season}: measurement_group="
                f"{df['measurement_group'].nunique()}，"
                f"预期={EXPECTED_GROUPS[season]}"
            )

        cross = (
            df.groupby("measurement_group")["Date"]
            .nunique()
            .gt(1)
            .sum()
        )
        if cross:
            raise ValueError(
                f"{season}: 有{cross}个measurement_group跨日期"
            )

        df["y_model"] = df[TARGET].map(RAW_TO_MODEL).astype(int)

        for d, g in df.groupby("Date", sort=True):
            if set(g["y_model"].unique()) != {0, 1, 2}:
                raise ValueError(
                    f"{season} {d.date()}: 单日不含完整三类"
                )

        df = df.reset_index(drop=True)
        df["row_id"] = np.arange(len(df), dtype=int)
        out[season] = df

        print(
            f"[数据核验] {season}: N={len(df)} | "
            f"days={df['Date'].nunique()} | "
            f"groups={df['measurement_group'].nunique()} | "
            f"classes={df[TARGET].value_counts().sort_index().to_dict()}"
        )

    return out


# =============================================================================
# 4. 读取12B已经冻结的XGB within-family配置
# =============================================================================

def load_12b_xgb_configs(path12b: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    required_sheets = [
        "04_Best_Params_by_Family",
        "06_common_family_inner",
        "07_Outer_Fold_Performance",
        "08_Outer_OOF_Summary",
        "10_Outer_Predictions",
    ]

    xls = pd.ExcelFile(path12b)
    missing_sheets = [s for s in required_sheets if s not in xls.sheet_names]
    if missing_sheets:
        raise ValueError(
            f"12B结果Excel缺少sheet：{missing_sheets}\n"
            f"实际：{xls.sheet_names}"
        )

    family_best = pd.read_excel(path12b, sheet_name="04_Best_Params_by_Family")
    common = pd.read_excel(path12b, sheet_name="06_common_family_inner")

    required_family_cols = [
        "Season",
        "Outer_fold",
        "Held_out_date",
        "Model",
        "Params_json",
        "Robust_Min",
        "Robust_Mean",
        "WorstDay_MacroF1",
        "Robust_WorstClassF1",
        "Stability_MaxSD",
        "Robust_BalancedAccuracy",
        "Robust_Kappa",
        "Within_family_selected",
    ]
    miss = [c for c in required_family_cols if c not in family_best.columns]
    if miss:
        raise ValueError(
            f"12B/04_Best_Params_by_Family缺少列：{miss}"
        )

    # 验证common family确实是XGB，并且来自inner-only排名
    required_common_cols = [
        "Model",
        "Final_common_family_rank",
        "Selected_for_Step12",
        "CrossSeason_Min_of_SeasonMean",
    ]
    miss2 = [c for c in required_common_cols if c not in common.columns]
    if miss2:
        raise ValueError(
            f"12B/06_common_family_inner缺少列：{miss2}"
        )

    common_sorted = common.sort_values("Final_common_family_rank")
    top = common_sorted.iloc[0]

    if str(top["Model"]).strip() != COMMON_FAMILY:
        raise RuntimeError(
            f"12B common family当前不是XGB，而是{top['Model']}。"
            "13-0拒绝继续，避免与12B冻结逻辑不一致。"
        )

    if not bool_from_excel(top["Selected_for_Step12"]):
        raise RuntimeError(
            "12B common family第一名没有Selected_for_Step12=True，"
            "请核查12B结果。"
        )

    # sheet04本身已经是每个family×fold的唯一selected配置，
    # 仍然额外检查Within_family_selected=True。
    xgb = family_best[
        (family_best["Model"].astype(str).str.strip() == COMMON_FAMILY)
        & family_best["Within_family_selected"].map(bool_from_excel)
    ].copy()

    if len(xgb) != 8:
        raise RuntimeError(
            f"应读取8个XGB within-family配置，实际={len(xgb)}。"
        )

    xgb["Held_out_date"] = pd.to_datetime(
        xgb["Held_out_date"]
    ).dt.normalize()

    # 每季必须4fold且日期唯一
    for season in ["Summer", "Winter"]:
        g = xgb[xgb["Season"] == season].copy()
        if len(g) != 4:
            raise RuntimeError(
                f"{season} XGB配置应为4折，实际={len(g)}"
            )
        if set(g["Outer_fold"].astype(int)) != {1, 2, 3, 4}:
            raise RuntimeError(
                f"{season} XGB Outer_fold不是1/2/3/4"
            )
        if g["Held_out_date"].nunique() != 4:
            raise RuntimeError(
                f"{season} XGB held-out date不唯一"
            )

    # Params_json必须可解析，且只允许XGB正式参数
    for _, r in xgb.iterrows():
        try:
            p = json.loads(r["Params_json"])
        except Exception as exc:
            raise ValueError(
                f"{r['Season']} F{r['Outer_fold']} Params_json无法解析"
            ) from exc

        required_params = {
            "selector_percentile",
            "weight_mode",
            "n_estimators",
            "learning_rate",
            "max_depth",
            "min_child_weight",
            "subsample",
            "colsample_bytree",
            "gamma",
            "reg_alpha",
            "reg_lambda",
        }
        missp = required_params - set(p)
        if missp:
            raise ValueError(
                f"{r['Season']} F{r['Outer_fold']} XGB参数缺少：{missp}"
            )

    xgb = xgb.sort_values(["Season", "Outer_fold"]).reset_index(drop=True)

    print(
        "[12B核验] common family = XGB | "
        f"cross-season minimum="
        f"{float(top['CrossSeason_Min_of_SeasonMean']):.6f}"
    )
    print("[12B核验] 已读取8个fold-specific XGB within-family配置；不重新调参。")

    return xgb, common_sorted


# =============================================================================
# 5. 预处理、XGB、指标 —— 与12B一致
# =============================================================================

def make_onehot():
    try:
        return OneHotEncoder(
            handle_unknown="ignore",
            sparse_output=False,
            dtype=np.float64,
        )
    except TypeError:
        return OneHotEncoder(
            handle_unknown="ignore",
            sparse=False,
            dtype=np.float64,
        )


def make_preprocessor():
    num_pipe = Pipeline([
        ("imputer", SimpleImputer(strategy="median"))
    ])
    cat_pipe = Pipeline([
        ("imputer", SimpleImputer(strategy="most_frequent")),
        ("onehot", make_onehot()),
    ])
    return ColumnTransformer(
        transformers=[
            ("num", num_pipe, EP_NUMERIC),
            ("cat", cat_pipe, EP_CATEGORICAL),
        ],
        remainder="drop",
        verbose_feature_names_out=False,
    )


def decode_xgb_config(raw_params: dict):
    raw = dict(raw_params)
    selector = int(raw.pop("selector_percentile"))
    weight_mode = str(raw.pop("weight_mode"))

    if selector not in {50, 65, 80, 100}:
        raise ValueError(f"非法selector_percentile={selector}")

    if weight_mode not in {"none", "balanced"}:
        raise ValueError(f"非法weight_mode={weight_mode}")

    return selector, weight_mode, raw


def make_xgb_estimator(params: dict, seed: int):
    p = dict(params)
    return XGBClassifier(
        objective="multi:softprob",
        num_class=3,
        eval_metric="mlogloss",
        n_estimators=int(p["n_estimators"]),
        learning_rate=float(p["learning_rate"]),
        max_depth=int(p["max_depth"]),
        min_child_weight=float(p["min_child_weight"]),
        subsample=float(p["subsample"]),
        colsample_bytree=float(p["colsample_bytree"]),
        gamma=float(p["gamma"]),
        reg_alpha=float(p["reg_alpha"]),
        reg_lambda=float(p["reg_lambda"]),
        random_state=seed,
        n_jobs=MODEL_N_JOBS,
        tree_method="hist",
    )


def make_xgb_pipeline(raw_config: dict, seed: int):
    selector, _, model_params = decode_xgb_config(raw_config)

    return Pipeline([
        ("preprocess", make_preprocessor()),
        (
            "selector",
            SelectPercentile(
                score_func=f_classif,
                percentile=selector,
            ),
        ),
        ("model", make_xgb_estimator(model_params, seed)),
    ])


def fit_pipeline(pipe, X, y, weight_mode: str):
    kwargs = {}
    if weight_mode == "balanced":
        kwargs["model__sample_weight"] = compute_sample_weight(
            "balanced", y=y
        )
    pipe.fit(X, y, **kwargs)
    return pipe


def metrics_bundle(y_true, y_pred):
    precision, recall, f1, support = precision_recall_fscore_support(
        y_true,
        y_pred,
        labels=CLASS_ORDER,
        zero_division=0,
    )

    out = {
        "Macro_F1": float(
            f1_score(
                y_true,
                y_pred,
                labels=CLASS_ORDER,
                average="macro",
                zero_division=0,
            )
        ),
        "Balanced_Accuracy": float(
            balanced_accuracy_score(y_true, y_pred)
        ),
        "Kappa": float(
            cohen_kappa_score(
                y_true, y_pred, labels=CLASS_ORDER
            )
        ),
        "Accuracy": float(
            accuracy_score(y_true, y_pred)
        ),
    }

    for i, cls in enumerate(CLASS_ORDER):
        raw_cls = MODEL_TO_RAW[cls]
        out[f"Class{raw_cls}_Precision"] = float(precision[i])
        out[f"Class{raw_cls}_Recall"] = float(recall[i])
        out[f"Class{raw_cls}_F1"] = float(f1[i])
        out[f"Class{raw_cls}_Support"] = int(support[i])

    out["Worst_Class_F1"] = float(np.min(f1))
    return out


def model_fit_seed(season: str, outer_fold: int) -> int:
    """
    与12B：
      model_fit_seed(season, fold, "XGB", phase_offset=900000)
    完全一致。
    """
    return (
        RANDOM_STATE
        + 20_000_000
        + OUTER_FIT_PHASE_OFFSET
        + (0 if season == "Summer" else 300000)
        + int(outer_fold) * 1000
        + CANDIDATE_MODELS_12B.index("XGB") * 31
    )


# =============================================================================
# 6. 特征名元数据 —— 供后续13A SHAP直接使用
# =============================================================================

def extract_selected_feature_metadata(pipe: Pipeline) -> pd.DataFrame:
    pre = pipe.named_steps["preprocess"]
    selector = pipe.named_steps["selector"]

    try:
        names = np.asarray(pre.get_feature_names_out(), dtype=object)
    except Exception:
        # 理论上现代sklearn可用；这里提供后备方案
        n_num = len(EP_NUMERIC)
        names_num = np.asarray(EP_NUMERIC, dtype=object)

        cat_names = []
        onehot = (
            pre.named_transformers_["cat"]
            .named_steps["onehot"]
        )
        cats = onehot.categories_
        for feature, levels in zip(EP_CATEGORICAL, cats):
            for level in levels:
                cat_names.append(f"{feature}_{level}")

        names = np.concatenate(
            [names_num, np.asarray(cat_names, dtype=object)]
        )

    support = selector.get_support()
    if len(names) != len(support):
        raise RuntimeError(
            f"预处理后特征数={len(names)}，selector mask={len(support)}，不一致"
        )

    rows = []
    selected_rank = 0
    for idx, (name, keep) in enumerate(zip(names, support)):
        name = str(name)

        # 供后续SHAP聚合到原变量：
        if name in EP_NUMERIC:
            parent = name
            kind = "numeric"
        elif name.startswith("Location_"):
            parent = "Location"
            kind = "onehot"
        elif name.startswith("Scene_"):
            parent = "Scene"
            kind = "onehot"
        else:
            # 防止新版sklearn产生不同命名
            if "Location" in name:
                parent = "Location"
                kind = "onehot"
            elif "Scene" in name:
                parent = "Scene"
                kind = "onehot"
            else:
                parent = name
                kind = "unknown"

        if keep:
            selected_rank += 1

        rows.append({
            "Transformed_index": idx,
            "Transformed_feature": name,
            "Parent_feature": parent,
            "Feature_kind": kind,
            "Selected": bool(keep),
            "Selected_order": selected_rank if keep else np.nan,
        })

    return pd.DataFrame(rows)


# =============================================================================
# 7. 断点续跑
# =============================================================================

def fold_key(season: str, fold: int) -> str:
    return f"{season}__F{int(fold)}"


def checkpoint_paths(run_dir: Path):
    checkpoint_dir = run_dir / "checkpoint"
    pred_dir = run_dir / "predictions"
    model_dir = run_dir / "models"
    feature_dir = run_dir / "feature_metadata"
    plot_ready_dir = run_dir / "plot_ready"

    for d in [checkpoint_dir, pred_dir, model_dir, feature_dir, plot_ready_dir]:
        d.mkdir(parents=True, exist_ok=True)

    return {
        "checkpoint_dir": checkpoint_dir,
        "checkpoint_csv": checkpoint_dir / "13-0_fold_results.csv",
        "pred_dir": pred_dir,
        "model_dir": model_dir,
        "feature_dir": feature_dir,
        "plot_ready_dir": plot_ready_dir,
    }


def fold_artifacts(paths: dict, season: str, fold: int):
    stem = f"13-0_{season}_F{fold}_XGB"
    return {
        "pred": paths["pred_dir"] / f"{stem}_predictions.csv",
        "model": paths["model_dir"] / f"{stem}.joblib",
        "features": paths["feature_dir"] / f"{stem}_features.csv",
    }


def load_checkpoint(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    df = pd.read_csv(path, encoding="utf-8-sig")
    return df


def checkpoint_fold_complete(
    ckpt: pd.DataFrame,
    paths: dict,
    season: str,
    fold: int,
) -> bool:
    if ckpt.empty:
        return False

    key = fold_key(season, fold)
    if "fold_key" not in ckpt.columns:
        return False

    if key not in set(ckpt["fold_key"].astype(str)):
        return False

    arts = fold_artifacts(paths, season, fold)
    return all(p.exists() for p in arts.values())


def append_checkpoint_row(
    checkpoint_csv: Path,
    row: dict,
):
    existing = load_checkpoint(checkpoint_csv)

    new = pd.DataFrame([row])

    if existing.empty:
        out = new
    else:
        existing = existing[
            existing["fold_key"].astype(str) != str(row["fold_key"])
        ]
        out = pd.concat([existing, new], ignore_index=True)

    out = out.sort_values(["Season", "Outer_fold"]).reset_index(drop=True)
    out.to_csv(
        checkpoint_csv,
        index=False,
        encoding="utf-8-sig",
    )


# =============================================================================
# 8. 配置一致性 / 防止不同数据混进同一断点目录
# =============================================================================

def build_run_config(input_file: Path, path12b: Path, smoke_test: bool):
    return {
        "script_version": SCRIPT_VERSION,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "purpose": (
            "Final strict cross-fitted predictive evaluation of the selected common XGB family; "
            "NO RETUNING / NO RESELECTION"
        ),
        "input_file": str(input_file.resolve()),
        "input_sha256": sha256_file(input_file),
        "source_12B_excel": str(path12b.resolve()),
        "source_12B_sha256": sha256_file(path12b),
        "common_family": COMMON_FAMILY,
        "parameter_source": "12B sheet 04_Best_Params_by_Family, XGB within-family selected",
        "feature_set": "EP",
        "EP_features": EP_FEATURES,
        "target": TARGET,
        "raw_to_model": RAW_TO_MODEL,
        "random_state": RANDOM_STATE,
        "outer_fit_phase_offset": OUTER_FIT_PHASE_OFFSET,
        "model_n_jobs": MODEL_N_JOBS,
        "expected_n": EXPECTED_N,
        "expected_groups": EXPECTED_GROUPS,
        "smoke_test": bool(smoke_test),
    }


def validate_or_write_config(
    config_path: Path,
    current: dict,
):
    if config_path.exists():
        old = json.loads(config_path.read_text(encoding="utf-8"))

        keys = [
            "input_sha256",
            "source_12B_sha256",
            "common_family",
            "feature_set",
            "target",
            "smoke_test",
        ]

        mismatch = {
            k: (old.get(k), current.get(k))
            for k in keys
            if old.get(k) != current.get(k)
        }

        if mismatch:
            raise RuntimeError(
                "检测到当前输入/12B来源与现有13-0断点目录不一致，"
                "为避免混合不同实验，程序拒绝续跑。\n"
                f"不一致项：{mismatch}\n"
                "如确认要从头开始，请使用 --reset。"
            )
    else:
        config_path.write_text(
            json.dumps(
                current,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )


# =============================================================================
# 9. 单fold严格XGB outer考试
# =============================================================================

def run_one_fold(
    data: dict[str, pd.DataFrame],
    config_row: pd.Series,
    paths: dict,
):
    season = str(config_row["Season"])
    fold = int(config_row["Outer_fold"])
    held = pd.Timestamp(config_row["Held_out_date"]).normalize()

    season_df = data[season]

    dates = sorted(season_df["Date"].unique())
    expected_held = pd.Timestamp(dates[fold - 1]).normalize()

    # 12B outer fold编号本来就是按日期排序生成；再次硬核验
    if held != expected_held:
        raise RuntimeError(
            f"{season} F{fold}: 12B记录held={held.date()}，"
            f"但源数据按日期排序F{fold}应为{expected_held.date()}。"
        )

    train = (
        season_df[season_df["Date"] != held]
        .copy()
        .reset_index(drop=True)
    )
    test = (
        season_df[season_df["Date"] == held]
        .copy()
        .reset_index(drop=True)
    )

    if len(test) == 0:
        raise RuntimeError(
            f"{season} F{fold}: held-out day没有测试样本"
        )

    train_groups = set(train["measurement_group"])
    test_groups = set(test["measurement_group"])
    overlap = train_groups & test_groups

    if overlap:
        raise RuntimeError(
            f"{season} F{fold}: Outer train/test存在"
            f"{len(overlap)}个measurement_group overlap"
        )

    raw_config = json.loads(config_row["Params_json"])
    selector, weight_mode, _ = decode_xgb_config(raw_config)
    seed = model_fit_seed(season, fold)

    print(
        f"\n[运行] {season} F{fold} | held={held.date()} | "
        f"train={len(train)} | test={len(test)} | "
        f"selector={selector}% | weight={weight_mode} | seed={seed}"
    )
    print(
        f"       12B inner证据: Robust_Min={float(config_row['Robust_Min']):.4f} | "
        f"Robust_Mean={float(config_row['Robust_Mean']):.4f} | "
        f"WorstDay={float(config_row['WorstDay_MacroF1']):.4f}"
    )

    t0 = time.time()

    pipe = make_xgb_pipeline(raw_config, seed)
    fit_pipeline(
        pipe,
        train[EP_FEATURES],
        train["y_model"].to_numpy(),
        weight_mode,
    )

    pred = pipe.predict(test[EP_FEATURES]).astype(int)
    proba = pipe.predict_proba(test[EP_FEATURES])

    ytrue = test["y_model"].to_numpy()
    metrics = metrics_bundle(ytrue, pred)

    if proba.shape != (len(test), 3):
        raise RuntimeError(
            f"{season} F{fold}: predict_proba shape={proba.shape}，应为(N,3)"
        )

    duration = time.time() - t0

    # 逐条OOF
    pred_df = pd.DataFrame({
        "Season": season,
        "Outer_fold": fold,
        "Held_out_date": held.strftime("%Y-%m-%d"),
        "row_id": test["row_id"].astype(int).to_numpy(),
        "Date": test["Date"].dt.strftime("%Y-%m-%d").to_numpy(),
        "y_true_model": ytrue.astype(int),
        "y_pred_model": pred.astype(int),
        "y_true_raw": [MODEL_TO_RAW[int(v)] for v in ytrue],
        "y_pred_raw": [MODEL_TO_RAW[int(v)] for v in pred],
        "P_Dissatisfied": proba[:, 0],
        "P_Neutral": proba[:, 1],
        "P_Satisfied": proba[:, 2],
        "Model": COMMON_FAMILY,
        "Selector_percentile": selector,
        "Weight_mode": weight_mode,
        "Fit_seed": seed,
    }).sort_values("row_id")

    # 特征筛选元数据
    feature_meta = extract_selected_feature_metadata(pipe)
    feature_meta.insert(0, "Season", season)
    feature_meta.insert(1, "Outer_fold", fold)
    feature_meta.insert(2, "Held_out_date", held.strftime("%Y-%m-%d"))

    # 保存fold artifacts
    arts = fold_artifacts(paths, season, fold)
    pred_df.to_csv(arts["pred"], index=False, encoding="utf-8-sig")
    feature_meta.to_csv(
        arts["features"],
        index=False,
        encoding="utf-8-sig",
    )
    joblib.dump(pipe, arts["model"])

    row = {
        "fold_key": fold_key(season, fold),
        "Season": season,
        "Outer_fold": fold,
        "Held_out_date": held.strftime("%Y-%m-%d"),
        "Model": COMMON_FAMILY,
        "N_train": len(train),
        "N_test": len(test),
        "Train_groups": len(train_groups),
        "Test_groups": len(test_groups),
        "Group_overlap": len(overlap),
        "Selector_percentile": selector,
        "Weight_mode": weight_mode,
        "Fit_seed": seed,
        "Robust_Min": float(config_row["Robust_Min"]),
        "Robust_Mean": float(config_row["Robust_Mean"]),
        "WorstDay_MacroF1_inner": float(config_row["WorstDay_MacroF1"]),
        "Robust_WorstClassF1_inner": float(config_row["Robust_WorstClassF1"]),
        "Stability_MaxSD_inner": float(config_row["Stability_MaxSD"]),
        "Robust_BalancedAccuracy_inner": float(config_row["Robust_BalancedAccuracy"]),
        "Robust_Kappa_inner": float(config_row["Robust_Kappa"]),
        "Params_json": config_row["Params_json"],
        "Selected_transformed_features": int(feature_meta["Selected"].sum()),
        "Fit_seconds": float(duration),
        "Macro_F1": metrics["Macro_F1"],
        "Balanced_Accuracy": metrics["Balanced_Accuracy"],
        "Kappa": metrics["Kappa"],
        "Accuracy": metrics["Accuracy"],
        "Worst_Class_F1": metrics["Worst_Class_F1"],
        "Class1_Precision": metrics["Class1_Precision"],
        "Class1_Recall": metrics["Class1_Recall"],
        "Class1_F1": metrics["Class1_F1"],
        "Class1_Support": metrics["Class1_Support"],
        "Class2_Precision": metrics["Class2_Precision"],
        "Class2_Recall": metrics["Class2_Recall"],
        "Class2_F1": metrics["Class2_F1"],
        "Class2_Support": metrics["Class2_Support"],
        "Class3_Precision": metrics["Class3_Precision"],
        "Class3_Recall": metrics["Class3_Recall"],
        "Class3_F1": metrics["Class3_F1"],
        "Class3_Support": metrics["Class3_Support"],
        "Completed_at": datetime.now().isoformat(timespec="seconds"),
    }

    append_checkpoint_row(paths["checkpoint_csv"], row)

    print(
        f"       Outer XGB Macro-F1={metrics['Macro_F1']:.4f} | "
        f"BalAcc={metrics['Balanced_Accuracy']:.4f} | "
        f"Kappa={metrics['Kappa']:.4f} | "
        f"WorstClass={metrics['Worst_Class_F1']:.4f} | "
        f"{fmt_seconds(duration)}"
    )

    return row


# =============================================================================
# 10. 汇总
# =============================================================================

def load_all_predictions(
    paths: dict,
    configs: pd.DataFrame,
) -> pd.DataFrame:
    dfs = []
    for _, r in configs.sort_values(["Season", "Outer_fold"]).iterrows():
        season = str(r["Season"])
        fold = int(r["Outer_fold"])
        p = fold_artifacts(paths, season, fold)["pred"]
        if not p.exists():
            raise FileNotFoundError(
                f"缺少fold预测文件：{p}"
            )
        dfs.append(pd.read_csv(p))

    out = pd.concat(dfs, ignore_index=True)
    out = out.sort_values(["Season", "row_id"]).reset_index(drop=True)
    return out


def build_oof_summary(
    pred_df: pd.DataFrame,
    fold_df: pd.DataFrame,
):
    summary_rows = []
    class_rows = []

    for season in ["Summer", "Winter"]:
        g = pred_df[pred_df["Season"] == season].copy()

        y = g["y_true_model"].astype(int).to_numpy()
        p = g["y_pred_model"].astype(int).to_numpy()
        m = metrics_bundle(y, p)

        ff = (
            fold_df[fold_df["Season"] == season]
            .copy()
            .sort_values("Outer_fold")
        )

        summary_rows.append({
            "Season": season,
            "N": len(g),
            "XGB_OOF_Macro_F1": m["Macro_F1"],
            "XGB_OOF_Balanced_Accuracy": m["Balanced_Accuracy"],
            "XGB_OOF_Kappa": m["Kappa"],
            "XGB_OOF_Accuracy": m["Accuracy"],
            "XGB_OOF_Worst_Class_F1": m["Worst_Class_F1"],
            "Dissatisfied_F1": m["Class1_F1"],
            "Neutral_F1": m["Class2_F1"],
            "Satisfied_F1": m["Class3_F1"],
            "Outer_fold_MacroF1_mean": float(ff["Macro_F1"].mean()),
            "Outer_fold_MacroF1_SD_descriptive": float(
                ff["Macro_F1"].std(ddof=1)
            ),
            "Outer_fold_MacroF1_worst": float(ff["Macro_F1"].min()),
            "Outer_fold_MacroF1_best": float(ff["Macro_F1"].max()),
            "Family": COMMON_FAMILY,
        })

        for cls in [1, 2, 3]:
            class_rows.append({
                "Season": season,
                "Class": cls,
                "Class_label": {
                    1: "Dissatisfied",
                    2: "Neutral",
                    3: "Satisfied",
                }[cls],
                "Precision": m[f"Class{cls}_Precision"],
                "Recall": m[f"Class{cls}_Recall"],
                "F1": m[f"Class{cls}_F1"],
                "Support": m[f"Class{cls}_Support"],
            })

    return pd.DataFrame(summary_rows), pd.DataFrame(class_rows)


def build_comparison_with_12b(
    path12b: Path,
    xgb_summary: pd.DataFrame,
):
    final12b = pd.read_excel(
        path12b,
        sheet_name="08_Outer_OOF_Summary",
    )

    rows = []
    for season in ["Summer", "Winter"]:
        a = xgb_summary[xgb_summary["Season"] == season].iloc[0]
        b = final12b[final12b["Season"] == season].iloc[0]

        rows.append({
            "Season": season,
            "12B_LocalWinnerPipeline_Macro_F1": float(b["OOF_Macro_F1"]),
            "13-0_FinalXGB_Macro_F1": float(a["XGB_OOF_Macro_F1"]),
            "Delta_FinalXGB_minus_12BLocalWinner": (
                float(a["XGB_OOF_Macro_F1"]) - float(b["OOF_Macro_F1"])
            ),
            "12B_LocalWinnerPipeline_BalAcc": float(b["OOF_Balanced_Accuracy"]),
            "13-0_FinalXGB_BalAcc": float(a["XGB_OOF_Balanced_Accuracy"]),
            "12B_LocalWinnerPipeline_Kappa": float(b["OOF_Kappa"]),
            "13-0_FinalXGB_Kappa": float(a["XGB_OOF_Kappa"]),
            "12B_LocalWinnerPipeline_Accuracy": float(b["OOF_Accuracy"]),
            "13-0_FinalXGB_Accuracy": float(a["XGB_OOF_Accuracy"]),
            "Interpretation": (
                "12B local-winner pipeline is retained only as a model-selection benchmark; "
                "13-0 provides the formal strict cross-fitted performance of the selected common XGB family."
            ),
        })

    return pd.DataFrame(rows)


def reproducibility_check_for_xgb_winner_folds(
    path12b: Path,
    pred13: pd.DataFrame,
    fold13: pd.DataFrame,
):
    """
    对12B本来就选择XGB作为Outer winner的fold做复现核验。
    在同一机器/版本下，13-0使用相同参数、相同seed、相同pipeline，
    理论上应复现这些fold的12B预测。
    不对非XGB winner fold作此要求。
    """
    outer12 = pd.read_excel(
        path12b,
        sheet_name="07_Outer_Fold_Performance",
    )
    pred12 = pd.read_excel(
        path12b,
        sheet_name="10_Outer_Predictions",
    )

    xgb_winners = outer12[
        outer12["Selected_model"].astype(str).str.strip() == "XGB"
    ].copy()

    rows = []

    for _, r in xgb_winners.iterrows():
        season = str(r["Season"])
        fold = int(r["Outer_fold"])

        a = pred12[
            (pred12["Season"] == season)
            & (pred12["Outer_fold"].astype(int) == fold)
        ][["row_id", "y_pred_model"]].copy()

        b = pred13[
            (pred13["Season"] == season)
            & (pred13["Outer_fold"].astype(int) == fold)
        ][["row_id", "y_pred_model"]].copy()

        merged = a.merge(
            b,
            on="row_id",
            suffixes=("_12B", "_13_0"),
            how="outer",
            indicator=True,
        )

        same_rows = bool((merged["_merge"] == "both").all())
        if same_rows:
            pred_match = float(
                (
                    merged["y_pred_model_12B"].astype(int)
                    == merged["y_pred_model_13_0"].astype(int)
                ).mean()
            )
        else:
            pred_match = np.nan

        f13 = fold13[
            (fold13["Season"] == season)
            & (fold13["Outer_fold"].astype(int) == fold)
        ].iloc[0]

        macro_diff = float(f13["Macro_F1"]) - float(r["Macro_F1"])

        rows.append({
            "Season": season,
            "Outer_fold": fold,
            "Held_out_date": pd.Timestamp(r["Held_out_date"]).strftime("%Y-%m-%d"),
            "12B_Selected_model": "XGB",
            "Prediction_row_set_identical": same_rows,
            "Prediction_exact_match_fraction": pred_match,
            "12B_Macro_F1": float(r["Macro_F1"]),
            "13-0_Macro_F1": float(f13["Macro_F1"]),
            "Macro_F1_difference": macro_diff,
            "Note": (
                "Expected to reproduce because 12B winner was already XGB; "
                "small discrepancies can occur after library/version changes."
            ),
        })

    return pd.DataFrame(rows)


def combine_feature_metadata(
    paths: dict,
    configs: pd.DataFrame,
):
    dfs = []
    for _, r in configs.iterrows():
        p = fold_artifacts(
            paths,
            str(r["Season"]),
            int(r["Outer_fold"]),
        )["features"]
        dfs.append(pd.read_csv(p))
    return pd.concat(dfs, ignore_index=True)


# =============================================================================
# 10B. 最新论文图表体系 plot-ready 输出
# =============================================================================

def export_plot_ready_v2(
    paths: dict,
    fold_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    class_df: pd.DataFrame,
    pred_df: pd.DataFrame,
):
    """输出Fig.4 / Fig.5 / Table4 / TableS2可直接读取的数据。"""
    out_dir = paths["plot_ready_dir"]

    # Fig.4：逐held-out day Macro-F1 + pooled OOF Macro-F1
    pooled_map = summary_df.set_index("Season")["XGB_OOF_Macro_F1"].to_dict()
    fig4 = fold_df[[
        "Season", "Outer_fold", "Held_out_date", "N_train", "N_test",
        "Macro_F1", "Balanced_Accuracy", "Kappa", "Accuracy", "Worst_Class_F1"
    ]].copy()
    fig4["Pooled_OOF_Macro_F1"] = fig4["Season"].map(pooled_map)
    fig4 = fig4.sort_values(["Season", "Outer_fold"]).reset_index(drop=True)
    fig4.to_csv(out_dir / "Fig4_DayToDay_Generalization_Ready.csv", index=False, encoding="utf-8-sig")

    # Fig.5：分类别性能
    fig5_class = class_df.copy().sort_values(["Season", "Class"]).reset_index(drop=True)
    fig5_class.to_csv(out_dir / "Fig5_ClassPerformance_Ready.csv", index=False, encoding="utf-8-sig")

    # Fig.5：row-normalized confusion matrix
    cm_rows = []
    labels = [0, 1, 2]
    label_names = {0: "Dissatisfied", 1: "Neutral", 2: "Satisfied"}
    for season in ["Summer", "Winter"]:
        g = pred_df[pred_df["Season"] == season].copy()
        y_true = g["y_true_model"].astype(int).to_numpy()
        y_pred = g["y_pred_model"].astype(int).to_numpy()
        cm_count = confusion_matrix(y_true, y_pred, labels=labels)
        row_sum = cm_count.sum(axis=1, keepdims=True)
        cm_norm = np.divide(cm_count, row_sum, out=np.zeros_like(cm_count, dtype=float), where=row_sum != 0)
        for i, true_cls in enumerate(labels):
            for j, pred_cls in enumerate(labels):
                cm_rows.append({
                    "Season": season,
                    "True_class_model": true_cls,
                    "True_class": label_names[true_cls],
                    "Predicted_class_model": pred_cls,
                    "Predicted_class": label_names[pred_cls],
                    "Count": int(cm_count[i, j]),
                    "Row_normalized": float(cm_norm[i, j]),
                    "Row_normalized_pct": float(cm_norm[i, j] * 100.0),
                })
    fig5_cm = pd.DataFrame(cm_rows)
    fig5_cm.to_csv(out_dir / "Fig5_ConfusionMatrix_Ready.csv", index=False, encoding="utf-8-sig")

    # Table 4：正文总体性能
    table4 = summary_df.copy()
    table4.to_csv(out_dir / "Table4_OverallPerformance_Ready.csv", index=False, encoding="utf-8-sig")

    # Table S2 Panel A：逐fold性能
    fold_df.sort_values(["Season", "Outer_fold"]).to_csv(
        out_dir / "TableS2_FoldPerformance_Ready.csv", index=False, encoding="utf-8-sig"
    )

    return fig4, fig5_class, fig5_cm, table4


# =============================================================================
# 11. Excel输出
# =============================================================================

def style_excel(path: Path):
    wb = load_workbook(path)

    fill = PatternFill("solid", fgColor="D4ECEF")
    header_font = Font(
        name="Times New Roman",
        size=10,
        bold=True,
    )
    body_font = Font(
        name="Times New Roman",
        size=10,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
                c.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                if isinstance(c.value, float):
                    c.number_format = "0.000"

        for idx, col_cells in enumerate(
            ws.columns, start=1
        ):
            max_len = 0
            for c in col_cells:
                if c.value is not None:
                    max_len = max(max_len, len(str(c.value)))
            ws.column_dimensions[
                get_column_letter(idx)
            ].width = min(max(max_len + 2, 10), 32)

    wb.save(path)


def export_excel(
    run_dir: Path,
    run_config: dict,
    xgb_configs: pd.DataFrame,
    fold_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    class_df: pd.DataFrame,
    pred_df: pd.DataFrame,
    comparison_df: pd.DataFrame,
    repro_df: pd.DataFrame,
    feature_meta: pd.DataFrame,
):
    output = run_dir / OUTPUT_XLSX_NAME

    config_table = pd.DataFrame(
        {
            "Item": [
                "Step",
                "Purpose",
                "Retuning",
                "Common family",
                "Input",
                "12B source",
                "Feature set",
                "Target",
                "Outer validation",
                "Parameter source",
                "Random state",
                "N Summer",
                "N Winter",
                "Groups Summer",
                "Groups Winter",
            ],
            "Value": [
                "13-0 XGB-only cross-fitted qualification",
                "Verify strict cross-day performance of the common interpretation family before SHAP",
                "NO",
                "XGB",
                run_config["input_file"],
                run_config["source_12B_excel"],
                "EP",
                TARGET,
                "4-fold leave-one-day-out separately by season",
                "12B sheet 04_Best_Params_by_Family / XGB within-family selected",
                RANDOM_STATE,
                EXPECTED_N["Summer"],
                EXPECTED_N["Winter"],
                EXPECTED_GROUPS["Summer"],
                EXPECTED_GROUPS["Winter"],
            ],
        }
    )

    param_cols = [
        "Season",
        "Outer_fold",
        "Held_out_date",
        "Model",
        "StageA_rank",
        "StageA_trial_number",
        "StageA_inner_Macro_F1",
        "FreshSGKF_MacroF1_mean",
        "FreshSGKF_MacroF1_SD",
        "InnerDayLODO_OOF_MacroF1",
        "WorstDay_MacroF1",
        "Day_MacroF1_SD",
        "Robust_Min",
        "Robust_Mean",
        "Robust_WorstClassF1",
        "Stability_MaxSD",
        "Robust_BalancedAccuracy",
        "Robust_Kappa",
        "Params_json",
    ]
    param_cols = [
        c for c in param_cols if c in xgb_configs.columns
    ]
    params_out = xgb_configs[param_cols].copy()

    # 日期格式统一
    if "Held_out_date" in params_out.columns:
        params_out["Held_out_date"] = pd.to_datetime(
            params_out["Held_out_date"]
        ).dt.strftime("%Y-%m-%d")

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        config_table.to_excel(
            writer,
            sheet_name="00_Configuration",
            index=False,
        )
        params_out.to_excel(
            writer,
            sheet_name="01_XGB_Config_Source",
            index=False,
        )
        fold_df.to_excel(
            writer,
            sheet_name="02_XGB_Outer_Folds",
            index=False,
        )
        summary_df.to_excel(
            writer,
            sheet_name="03_XGB_OOF_Summary",
            index=False,
        )
        class_df.to_excel(
            writer,
            sheet_name="04_XGB_Class_OOF",
            index=False,
        )
        pred_df.to_excel(
            writer,
            sheet_name="05_XGB_Predictions",
            index=False,
        )
        comparison_df.to_excel(
            writer,
            sheet_name="06_12B_LocalWinner_Comparison",
            index=False,
        )
        repro_df.to_excel(
            writer,
            sheet_name="07_12B_XGB_Reproduction",
            index=False,
        )
        feature_meta.to_excel(
            writer,
            sheet_name="08_Feature_Selection",
            index=False,
        )

    style_excel(output)
    return output


# =============================================================================
# 12. 路径解析
# =============================================================================

def resolve_paths(args):
    base_dir = (
        Path(args.base_dir).expanduser().resolve()
        if args.base_dir
        else DEFAULT_BASE_DIR
    )

    input_file = (
        Path(args.input).expanduser().resolve()
        if args.input
        else base_dir / INPUT_NAME
    )

    dir12b = base_dir / DIR_12B_NAME
    path12b = (
        Path(args.result12b).expanduser().resolve()
        if args.result12b
        else resolve_xlsx(
            dir12b / FILE_12B_NAME,
            dir12b,
            "12B",
        )
    )

    suffix = "_SMOKE_TEST" if args.smoke_test else ""
    run_dir = base_dir / f"{RUN_DIR_NAME}{suffix}"

    if args.reset and run_dir.exists():
        print(f"[RESET] 删除旧输出目录：{run_dir}")
        shutil.rmtree(run_dir)

    run_dir.mkdir(parents=True, exist_ok=True)

    if not input_file.exists():
        raise FileNotFoundError(
            f"找不到真实源数据：\n{input_file}"
        )
    if not path12b.exists():
        raise FileNotFoundError(
            f"找不到12B结果Excel：\n{path12b}"
        )

    return base_dir, input_file, path12b, run_dir


# =============================================================================
# 13. 主流程
# =============================================================================

def main():
    parser = argparse.ArgumentParser(
        description=(
            "13-0：固定XGB common family，使用12B既有inner-selected参数，"
            "执行严格cross-fitted LODO性能验证；不重新调参。"
        )
    )
    parser.add_argument(
        "--base-dir",
        default=None,
        help="data_analyse目录；默认脚本所在目录",
    )
    parser.add_argument(
        "--input",
        default=None,
        help="真实源数据xlsx；默认自动读取",
    )
    parser.add_argument(
        "--result12b",
        default=None,
        help="12B正式结果xlsx；默认自动读取",
    )
    parser.add_argument(
        "--smoke-test",
        action="store_true",
        help="只跑Summer F1，独立SMOKE_TEST目录，不可用于论文",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="删除对应13-0输出目录后从头重跑",
    )

    args = parser.parse_args()

    base_dir, input_file, path12b, run_dir = resolve_paths(args)

    # 日志
    log_path = run_dir / LOG_NAME
    log_file = log_path.open("a", encoding="utf-8", buffering=1)
    original_stdout = sys.stdout
    original_stderr = sys.stderr
    sys.stdout = Tee(original_stdout, log_file)
    sys.stderr = Tee(original_stderr, log_file)

    try:
        print("\n" + "=" * 108)
        print("13-0 V2：XGB最终共同模型 / 严格Cross-Fitted LODO 预测评估")
        print("=" * 108)
        print(f"开始：{datetime.now()}")
        print(f"Base dir：{base_dir}")
        print(f"输入数据：{input_file}")
        print(f"12B结果：{path12b}")
        print(f"输出目录：{run_dir}")
        print(f"Smoke test：{args.smoke_test}")
        print("重新调参：False")
        print("参数来源：12B既有XGB within-family inner/stability选择结果")
        print("说明：13-0 V2作为选定XGB后的正式cross-fitted预测评估；12B local-winner仅作为筛选参考。")
        print("=" * 108)

        # 防止不同版本混用
        run_config = build_run_config(
            input_file,
            path12b,
            args.smoke_test,
        )
        config_path = run_dir / CONFIG_JSON_NAME
        validate_or_write_config(
            config_path,
            run_config,
        )

        # 读取
        data = load_data(input_file)
        xgb_configs, common_table = load_12b_xgb_configs(path12b)

        # 核验held dates与源数据完全一致
        for season in ["Summer", "Winter"]:
            dates_data = [
                pd.Timestamp(x).normalize()
                for x in sorted(data[season]["Date"].unique())
            ]
            dates_cfg = [
                pd.Timestamp(x).normalize()
                for x in (
                    xgb_configs[xgb_configs["Season"] == season]
                    .sort_values("Outer_fold")["Held_out_date"]
                    .tolist()
                )
            ]
            if dates_data != dates_cfg:
                raise RuntimeError(
                    f"{season}: 12B XGB配置日期与源数据4个日期不一致。\n"
                    f"数据={dates_data}\n12B={dates_cfg}"
                )

        if args.smoke_test:
            configs_run = xgb_configs[
                (xgb_configs["Season"] == "Summer")
                & (xgb_configs["Outer_fold"].astype(int) == 1)
            ].copy()
        else:
            configs_run = xgb_configs.copy()

        paths = checkpoint_paths(run_dir)
        ckpt = load_checkpoint(paths["checkpoint_csv"])

        completed = 0
        for _, r in configs_run.iterrows():
            if checkpoint_fold_complete(
                ckpt,
                paths,
                str(r["Season"]),
                int(r["Outer_fold"]),
            ):
                completed += 1

        total = len(configs_run)

        print(
            f"\n[运行计划] folds={total} | "
            f"断点已完成={completed} | "
            f"本次待跑={total-completed}"
        )
        print_progress(completed, total, time.time())

        start = time.time()

        # 逐fold执行
        for _, r in configs_run.sort_values(
            ["Season", "Outer_fold"]
        ).iterrows():
            season = str(r["Season"])
            fold = int(r["Outer_fold"])

            ckpt = load_checkpoint(paths["checkpoint_csv"])

            if checkpoint_fold_complete(
                ckpt,
                paths,
                season,
                fold,
            ):
                print(
                    f"\n[RESUME跳过] {season} F{fold} 已完整完成"
                )
                continue

            run_one_fold(
                data=data,
                config_row=r,
                paths=paths,
            )

            ckpt = load_checkpoint(paths["checkpoint_csv"])
            done_now = sum(
                checkpoint_fold_complete(
                    ckpt,
                    paths,
                    str(rr["Season"]),
                    int(rr["Outer_fold"]),
                )
                for _, rr in configs_run.iterrows()
            )
            print_progress(done_now, total, start)

        # smoke只验证链条，不生成正式两季汇总
        if args.smoke_test:
            print("\n" + "=" * 108)
            print("13-0 SMOKE TEST完成：只验证Summer F1链条，不可用于论文。")
            print("正式运行请去掉 --smoke-test。")
            print(f"模型：{paths['model_dir']}")
            print(f"预测：{paths['pred_dir']}")
            print(f"结束：{datetime.now()}")
            print("=" * 108)
            return

        # 正式必须8折齐全
        ckpt = load_checkpoint(paths["checkpoint_csv"])
        done = sum(
            checkpoint_fold_complete(
                ckpt,
                paths,
                str(r["Season"]),
                int(r["Outer_fold"]),
            )
            for _, r in xgb_configs.iterrows()
        )
        if done != 8:
            raise RuntimeError(
                f"正式13-0汇总要求8/8 fold完成，当前={done}/8"
            )

        fold_df = (
            ckpt.sort_values(["Season", "Outer_fold"])
            .reset_index(drop=True)
        )

        pred_df = load_all_predictions(
            paths,
            xgb_configs,
        )

        # 1001条OOF硬核验
        if len(pred_df) != 1001:
            raise RuntimeError(
                f"XGB OOF逐条预测应为1001，实际={len(pred_df)}"
            )
        if int((pred_df["Season"] == "Summer").sum()) != 477:
            raise RuntimeError("Summer XGB OOF不是477条")
        if int((pred_df["Season"] == "Winter").sum()) != 524:
            raise RuntimeError("Winter XGB OOF不是524条")

        # 每个原始row_id每季只出现一次
        for season, n in EXPECTED_N.items():
            g = pred_df[pred_df["Season"] == season]
            if g["row_id"].nunique() != n:
                raise RuntimeError(
                    f"{season}: XGB OOF row_id没有做到每人恰好一次"
                )

        summary_df, class_df = build_oof_summary(
            pred_df,
            fold_df,
        )

        comparison_df = build_comparison_with_12b(
            path12b,
            summary_df,
        )

        repro_df = reproducibility_check_for_xgb_winner_folds(
            path12b,
            pred_df,
            fold_df,
        )

        feature_meta = combine_feature_metadata(
            paths,
            xgb_configs,
        )

        # 为最新论文图表体系生成直接可用数据
        export_plot_ready_v2(
            paths=paths,
            fold_df=fold_df,
            summary_df=summary_df,
            class_df=class_df,
            pred_df=pred_df,
        )

        output_excel = export_excel(
            run_dir=run_dir,
            run_config=run_config,
            xgb_configs=xgb_configs,
            fold_df=fold_df,
            summary_df=summary_df,
            class_df=class_df,
            pred_df=pred_df,
            comparison_df=comparison_df,
            repro_df=repro_df,
            feature_meta=feature_meta,
        )

        # 另外输出主表CSV，方便13A与论文制表
        summary_df.to_csv(
            run_dir / "Table4_OverallPerformance_Ready.csv",
            index=False,
            encoding="utf-8-sig",
        )
        class_df.to_csv(
            run_dir / "Fig5_ClassPerformance_Ready.csv",
            index=False,
            encoding="utf-8-sig",
        )

        print("\n" + "=" * 108)
        print("13-0 V2正式完成：selected common XGB family严格Cross-Fitted LODO预测评估")
        print("=" * 108)

        for _, r in summary_df.iterrows():
            print(
                f"{r['Season']}: "
                f"Macro-F1={r['XGB_OOF_Macro_F1']:.4f} | "
                f"BalAcc={r['XGB_OOF_Balanced_Accuracy']:.4f} | "
                f"Kappa={r['XGB_OOF_Kappa']:.4f} | "
                f"Accuracy={r['XGB_OOF_Accuracy']:.4f} | "
                f"Class F1="
                f"{r['Dissatisfied_F1']:.4f}/"
                f"{r['Neutral_F1']:.4f}/"
                f"{r['Satisfied_F1']:.4f} | "
                f"Fold SD={r['Outer_fold_MacroF1_SD_descriptive']:.4f}"
            )

        print("\n与12B最终pipeline比较：")
        for _, r in comparison_df.iterrows():
            print(
                f"{r['Season']}: 12B={r['12B_LocalWinnerPipeline_Macro_F1']:.4f} "
                f"| XGB-only={r['13-0_FinalXGB_Macro_F1']:.4f} "
                f"| Δ={r['Delta_FinalXGB_minus_12BLocalWinner']:+.4f}"
            )

        if len(repro_df):
            print("\n12B原本就是XGB winner的fold复现核验：")
            for _, r in repro_df.iterrows():
                print(
                    f"{r['Season']} F{int(r['Outer_fold'])}: "
                    f"prediction exact match="
                    f"{r['Prediction_exact_match_fraction']:.3f} | "
                    f"Macro-F1 diff={r['Macro_F1_difference']:+.8f}"
                )

        print(f"\n结果Excel：{output_excel}")
        print(f"8个XGB模型：{paths['model_dir']}")
        print(f"OOF预测：{paths['pred_dir']}")
        print(f"特征元数据：{paths['feature_dir']}")
        print(
            "\n下一步13A：直接读取这8个XGB pipeline，"
            "分别对各自held-out day计算cross-fitted SHAP；"
            "不再重新调参。"
        )
        print(f"结束：{datetime.now()}")
        print("=" * 108)

    finally:
        sys.stdout = original_stdout
        sys.stderr = original_stderr
        log_file.close()


if __name__ == "__main__":
    main()
