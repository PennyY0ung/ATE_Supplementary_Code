from __future__ import annotations

import argparse
import hashlib
import json
import os
import shutil
import sys
import time
import warnings
from datetime import datetime, timedelta
from pathlib import Path
from typing import Any

import joblib
import numpy as np
import pandas as pd

from sklearn.base import clone
from sklearn.compose import ColumnTransformer
from sklearn.feature_selection import SelectPercentile, f_classif
from sklearn.impute import SimpleImputer
from sklearn.metrics import (
    accuracy_score,
    balanced_accuracy_score,
    cohen_kappa_score,
    confusion_matrix,
    f1_score,
    precision_recall_fscore_support,
)
from sklearn.pipeline import Pipeline
from sklearn.preprocessing import OneHotEncoder
from sklearn.utils.class_weight import compute_sample_weight

try:
    from xgboost import XGBClassifier
except Exception as exc:
    raise ImportError(
        "Final XGB evaluation requires xgboost. Install with: pip install xgboost"
    ) from exc

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except Exception as exc:
    raise ImportError("Excel export requires openpyxl. Install with: pip install openpyxl") from exc


# =============================================================================
# Fixed paths and analysis definitions
# =============================================================================

SCRIPT_VERSION = "13-0_v2_final_20260821"

DEFAULT_BASE_DIR = Path(__file__).resolve().parent

INPUT_NAME = "Full_Analytical_Dataset.xlsx"

DIR_12B_NAME = (
    "Hyperparameter_Optimization_Results"
)
FILE_12B_NAME = (
    "Hyperparameter_Optimization_Results.xlsx"
)

RUN_DIR_NAME = "Final_XGB_LODO_Results"
OUTPUT_XLSX_NAME = "Final_XGB_LODO_Results.xlsx"
CONFIG_JSON_NAME = "Final_XGB_LODO_Config.json"
LOG_NAME = "Final_XGB_LODO_Log.log"

SHEET_MAP = {"Summer": "夏季", "Winter": "冬季"}
EXPECTED_N = {"Summer": 477, "Winter": 524}
EXPECTED_GROUPS = {"Summer": 190, "Winter": 296}

TARGET = "Thermal_satisfaction_3"
RAW_TO_MODEL = {1: 0, 2: 1, 3: 2}
MODEL_TO_RAW = {0: 1, 1: 2, 2: 3}
CLASS_ORDER = [0, 1, 2]

OBJECTIVE = [
    "Temperature",
    "Wind_speed",
    "Humidity",
    "Globe_temperature",
    "Illumination",
]
SPATIAL = ["Location", "Scene"]
SUBJECTIVE = [
    "TSV", "TCV", "TPV", "TAV",
    "SSV", "SCV", "SPV", "SAV",
    "HSV", "WSV", "HPV", "WPV",
]
EP_NUMERIC = OBJECTIVE + SUBJECTIVE
EP_CATEGORICAL = SPATIAL
EP_FEATURES = EP_NUMERIC + EP_CATEGORICAL

COMMON_FAMILY = "XGB"

# Must agree with the five-family frozen model ordering.
CANDIDATE_MODELS_12B = ["GB", "CAT", "ET", "LGBM", "XGB"]
RANDOM_STATE = 20260816
MODEL_N_JOBS = -1
OUTER_FIT_PHASE_OFFSET = 900000


# =============================================================================
# Logging, progress, and utilities
# =============================================================================

class Tee:
    def __init__(self, *streams):
        self.streams = streams

    def write(self, data):
        for s in self.streams:
            s.write(data)
            s.flush()

    def flush(self):
        for s in self.streams:
            s.flush()


def fmt_seconds(seconds: float) -> str:
    if not np.isfinite(seconds):
        return "--"
    seconds = max(float(seconds), 0.0)
    if seconds < 60:
        return f"{seconds:.0f}s"
    if seconds < 3600:
        m, s = divmod(int(seconds), 60)
        return f"{m}m{s:02d}s"
    h, rem = divmod(int(seconds), 3600)
    m, s = divmod(rem, 60)
    return f"{h}h{m:02d}m{s:02d}s"


def ascii_bar(done: int, total: int, width: int = 30) -> str:
    if total <= 0:
        return "[" + "." * width + "]"
    frac = min(max(done / total, 0), 1)
    n = int(round(frac * width))
    if done >= total:
        return "[" + "=" * width + "]"
    n = min(n, width - 1)
    return "[" + "=" * n + ">" + "." * (width - n - 1) + "]"


def print_progress(done: int, total: int, start_time: float):
    elapsed = time.time() - start_time
    if done > 0:
        eta = elapsed / done * max(total - done, 0)
        finish = datetime.now() + timedelta(seconds=eta)
        finish_txt = finish.strftime("%Y-%m-%d %H:%M:%S")
    else:
        eta = float("nan")
        finish_txt = "--"

    print(
        f"[Overall progress] {ascii_bar(done, total)} "
        f"{done}/{total} folds ({done/total*100 if total else 0:.1f}%) "
        f"| Elapsed {fmt_seconds(elapsed)} "
        f"| ETA {fmt_seconds(eta)} "
        f"| Estimated completion {finish_txt}"
    )


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    h = hashlib.sha256()
    with path.open("rb") as f:
        while True:
            chunk = f.read(chunk_size)
            if not chunk:
                break
            h.update(chunk)
    return h.hexdigest()


def bool_from_excel(v: Any) -> bool:
    if isinstance(v, (bool, np.bool_)):
        return bool(v)
    s = str(v).strip().lower()
    return s in {"true", "1", "yes", "y"}


def resolve_xlsx(exact: Path, folder: Path, keyword: str) -> Path:
    if exact.exists():
        return exact

    if not folder.exists():
        raise FileNotFoundError(f"Directory not found: {folder}")

    candidates = [
        p for p in folder.glob("*.xlsx")
        if keyword.lower() in p.name.lower() and not p.name.startswith("~$")
    ]
    if len(candidates) == 1:
        return candidates[0]

    all_xlsx = [p for p in folder.glob("*.xlsx") if not p.name.startswith("~$")]
    if len(all_xlsx) == 1:
        return all_xlsx[0]

    raise FileNotFoundError(
        f"Cannot unambiguously identify {keyword} results workbook.\n"
        f"Directory: {folder}\n"
        f"Candidates: {[p.name for p in all_xlsx]}"
    )


# =============================================================================
# Load data with the same strict checks as Stage 12B.
# =============================================================================

def load_data(input_file: Path) -> dict[str, pd.DataFrame]:
    raw = pd.read_excel(
        input_file,
        sheet_name=list(SHEET_MAP.values()),
        engine="openpyxl",
    )

    required = ["Date", "measurement_group", TARGET] + EP_FEATURES
    out: dict[str, pd.DataFrame] = {}

    for season, sheet in SHEET_MAP.items():
        df = raw[sheet].copy()
        df.columns = [str(c).strip() for c in df.columns]

        missing = [c for c in required if c not in df.columns]
        if missing:
            raise KeyError(f"{season} is missing columns: {missing}")

        if len(df) != EXPECTED_N[season]:
            raise ValueError(
                f"{season}: N={len(df)}, expected={EXPECTED_N[season]}"
            )

        df["Date"] = pd.to_datetime(
            df["Date"], errors="coerce"
        ).dt.normalize()

        df[TARGET] = pd.to_numeric(df[TARGET], errors="coerce")

        for c in EP_NUMERIC:
            df[c] = pd.to_numeric(df[c], errors="coerce")

        for c in EP_CATEGORICAL:
            df[c] = df[c].astype(str).str.strip()

        df["measurement_group"] = (
            df["measurement_group"].astype(str).str.strip()
        )

        if df[required].isna().any().any():
            bad = df[required].isna().sum()
            raise ValueError(
                f"{season} contains missing or invalid values: "
                f"{bad[bad > 0].to_dict()}"
            )

        df[TARGET] = df[TARGET].astype(int)

        if set(df[TARGET].unique()) != {1, 2, 3}:
            raise ValueError(
                f"{season}: target must include all three classes (1/2/3)"
            )

        if df["Date"].nunique() != 4:
            raise ValueError(
                f"{season}: survey days={df['Date'].nunique()}, expected four"
            )

        if df["measurement_group"].nunique() != EXPECTED_GROUPS[season]:
            raise ValueError(
                f"{season}: measurement_group="
                f"{df['measurement_group'].nunique()},"
                f"expected={EXPECTED_GROUPS[season]}"
            )

        cross = (
            df.groupby("measurement_group")["Date"]
            .nunique()
            .gt(1)
            .sum()
        )
        if cross:
            raise ValueError(
                f"{season}: {cross} measurement groups span multiple survey dates"
            )

        df["y_model"] = df[TARGET].map(RAW_TO_MODEL).astype(int)

        for d, g in df.groupby("Date", sort=True):
            if set(g["y_model"].unique()) != {0, 1, 2}:
                raise ValueError(
                    f"{season} {d.date()}: day lacks one or more outcome classes"
                )

        df = df.reset_index(drop=True)
        df["row_id"] = np.arange(len(df), dtype=int)
        out[season] = df

        print(
            f"[Data validation] {season}: N={len(df)} | "
            f"days={df['Date'].nunique()} | "
            f"groups={df['measurement_group'].nunique()} | "
            f"classes={df[TARGET].value_counts().sort_index().to_dict()}"
        )

    return out


# =============================================================================
# Read frozen XGB configurations from Stage 12B results
# =============================================================================

def load_12b_xgb_configs(path12b: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    required_sheets = [
        "04_Best_Family_Configs",
        "06_Common_Family_Selection",
        "07_Outer_Fold_Performance",
        "08_Pooled_OOF_Summary",
        "10_Outer_Predictions",
    ]

    xls = pd.ExcelFile(path12b)
    missing_sheets = [s for s in required_sheets if s not in xls.sheet_names]
    if missing_sheets:
        raise ValueError(
            f"Stage 12B workbook is missing sheets: {missing_sheets}\n"
            f"Actual: {xls.sheet_names}"
        )

    family_best = pd.read_excel(path12b, sheet_name="04_Best_Family_Configs")
    common = pd.read_excel(path12b, sheet_name="06_Common_Family_Selection")

    required_family_cols = [
        "Season",
        "Outer_fold",
        "Held_out_date",
        "Model",
        "Params_json",
        "Robust_Min",
        "Robust_Mean",
        "WorstDay_MacroF1",
        "Robust_WorstClassF1",
        "Stability_MaxSD",
        "Robust_BalancedAccuracy",
        "Robust_Kappa",
        "Within_family_selected",
    ]
    miss = [c for c in required_family_cols if c not in family_best.columns]
    if miss:
        raise ValueError(
            f"Stage 12B/04_Best_Family_Configs missing columns: {miss}"
        )

    # Validate the common-family selection and inner-evidence ranking.
    required_common_cols = [
        "Model",
        "Final_common_family_rank",
        "Selected_for_Step12",
        "CrossSeason_Min_of_SeasonMean",
    ]
    miss2 = [c for c in required_common_cols if c not in common.columns]
    if miss2:
        raise ValueError(
            f"Stage 12B/06_Common_Family_Selection missing columns: {miss2}"
        )

    common_sorted = common.sort_values("Final_common_family_rank")
    top = common_sorted.iloc[0]

    if str(top["Model"]).strip() != COMMON_FAMILY:
        raise RuntimeError(
            f"Stage 12B common family is not XGB; found {top['Model']}."
            "Final evaluation cannot proceed because the frozen Stage 12B family differs."
        )

    if not bool_from_excel(top["Selected_for_Step12"]):
        raise RuntimeError(
            "The top common-family row is not marked Selected_for_Step12=True; "
            "check the Stage 12B results."
        )

    # Sheet 04 contains one selected configuration per family and fold.
    # Also verify the Within_family_selected flag explicitly.
    xgb = family_best[
        (family_best["Model"].astype(str).str.strip() == COMMON_FAMILY)
        & family_best["Within_family_selected"].map(bool_from_excel)
    ].copy()

    if len(xgb) != 8:
        raise RuntimeError(
            f"Expected eight within-family XGB configurations; found={len(xgb)}."
        )

    xgb["Held_out_date"] = pd.to_datetime(
        xgb["Held_out_date"]
    ).dt.normalize()

    # Each season must have four distinct folds and held-out dates.
    for season in ["Summer", "Winter"]:
        g = xgb[xgb["Season"] == season].copy()
        if len(g) != 4:
            raise RuntimeError(
                f"{season} XGB should have four fold configurations; found={len(g)}"
            )
        if set(g["Outer_fold"].astype(int)) != {1, 2, 3, 4}:
            raise RuntimeError(
                f"{season} XGB folds must be numbered 1/2/3/4"
            )
        if g["Held_out_date"].nunique() != 4:
            raise RuntimeError(
                f"{season} XGB held-out dates are not unique"
            )

    # Parse all configuration JSON values and check required XGB fields.
    for _, r in xgb.iterrows():
        try:
            p = json.loads(r["Params_json"])
        except Exception as exc:
            raise ValueError(
                f"{r['Season']} F{r['Outer_fold']} Params_json cannot be parsed"
            ) from exc

        required_params = {
            "selector_percentile",
            "weight_mode",
            "n_estimators",
            "learning_rate",
            "max_depth",
            "min_child_weight",
            "subsample",
            "colsample_bytree",
            "gamma",
            "reg_alpha",
            "reg_lambda",
        }
        missp = required_params - set(p)
        if missp:
            raise ValueError(
                f"{r['Season']} F{r['Outer_fold']} XGB parameters are missing: {missp}"
            )

    xgb = xgb.sort_values(["Season", "Outer_fold"]).reset_index(drop=True)

    print(
        "[Stage 12B Check] common family = XGB | "
        f"cross-season minimum="
        f"{float(top['CrossSeason_Min_of_SeasonMean']):.6f}"
    )
    print("[Stage 12B Check] Loaded eight frozen fold-specific XGB configurations; no retuning.")

    return xgb, common_sorted


# =============================================================================
# Preprocessing, XGB, and metrics retained from Stage 12B.
# =============================================================================

def make_onehot():
    try:
        return OneHotEncoder(
            handle_unknown="ignore",
            sparse_output=False,
            dtype=np.float64,
        )
    except TypeError:
        return OneHotEncoder(
            handle_unknown="ignore",
            sparse=False,
            dtype=np.float64,
        )


def make_preprocessor():
    num_pipe = Pipeline([
        ("imputer", SimpleImputer(strategy="median"))
    ])
    cat_pipe = Pipeline([
        ("imputer", SimpleImputer(strategy="most_frequent")),
        ("onehot", make_onehot()),
    ])
    return ColumnTransformer(
        transformers=[
            ("num", num_pipe, EP_NUMERIC),
            ("cat", cat_pipe, EP_CATEGORICAL),
        ],
        remainder="drop",
        verbose_feature_names_out=False,
    )


def decode_xgb_config(raw_params: dict):
    raw = dict(raw_params)
    selector = int(raw.pop("selector_percentile"))
    weight_mode = str(raw.pop("weight_mode"))

    if selector not in {50, 65, 80, 100}:
        raise ValueError(f"Invalid selector_percentile={selector}")

    if weight_mode not in {"none", "balanced"}:
        raise ValueError(f"Invalid weight_mode={weight_mode}")

    return selector, weight_mode, raw


def make_xgb_estimator(params: dict, seed: int):
    p = dict(params)
    return XGBClassifier(
        objective="multi:softprob",
        num_class=3,
        eval_metric="mlogloss",
        n_estimators=int(p["n_estimators"]),
        learning_rate=float(p["learning_rate"]),
        max_depth=int(p["max_depth"]),
        min_child_weight=float(p["min_child_weight"]),
        subsample=float(p["subsample"]),
        colsample_bytree=float(p["colsample_bytree"]),
        gamma=float(p["gamma"]),
        reg_alpha=float(p["reg_alpha"]),
        reg_lambda=float(p["reg_lambda"]),
        random_state=seed,
        n_jobs=MODEL_N_JOBS,
        tree_method="hist",
    )


def make_xgb_pipeline(raw_config: dict, seed: int):
    selector, _, model_params = decode_xgb_config(raw_config)

    return Pipeline([
        ("preprocess", make_preprocessor()),
        (
            "selector",
            SelectPercentile(
                score_func=f_classif,
                percentile=selector,
            ),
        ),
        ("model", make_xgb_estimator(model_params, seed)),
    ])


def fit_pipeline(pipe, X, y, weight_mode: str):
    kwargs = {}
    if weight_mode == "balanced":
        kwargs["model__sample_weight"] = compute_sample_weight(
            "balanced", y=y
        )
    pipe.fit(X, y, **kwargs)
    return pipe


def metrics_bundle(y_true, y_pred):
    precision, recall, f1, support = precision_recall_fscore_support(
        y_true,
        y_pred,
        labels=CLASS_ORDER,
        zero_division=0,
    )

    out = {
        "Macro_F1": float(
            f1_score(
                y_true,
                y_pred,
                labels=CLASS_ORDER,
                average="macro",
                zero_division=0,
            )
        ),
        "Balanced_Accuracy": float(
            balanced_accuracy_score(y_true, y_pred)
        ),
        "Kappa": float(
            cohen_kappa_score(
                y_true, y_pred, labels=CLASS_ORDER
            )
        ),
        "Accuracy": float(
            accuracy_score(y_true, y_pred)
        ),
    }

    for i, cls in enumerate(CLASS_ORDER):
        raw_cls = MODEL_TO_RAW[cls]
        out[f"Class{raw_cls}_Precision"] = float(precision[i])
        out[f"Class{raw_cls}_Recall"] = float(recall[i])
        out[f"Class{raw_cls}_F1"] = float(f1[i])
        out[f"Class{raw_cls}_Support"] = int(support[i])

    out["Worst_Class_F1"] = float(np.min(f1))
    return out


def model_fit_seed(season: str, outer_fold: int) -> int:
    """Matches Stage 12B model_fit_seed(season, fold, 'XGB', phase_offset=900000)."""
    return (
        RANDOM_STATE
        + 20_000_000
        + OUTER_FIT_PHASE_OFFSET
        + (0 if season == "Summer" else 300000)
        + int(outer_fold) * 1000
        + CANDIDATE_MODELS_12B.index("XGB") * 31
    )


# =============================================================================
# Extract feature metadata for subsequent cross-fitted SHAP calculations.
# =============================================================================

def extract_selected_feature_metadata(pipe: Pipeline) -> pd.DataFrame:
    pre = pipe.named_steps["preprocess"]
    selector = pipe.named_steps["selector"]

    try:
        names = np.asarray(pre.get_feature_names_out(), dtype=object)
    except Exception:
        # Fallback for environments lacking get_feature_names_out support.
        n_num = len(EP_NUMERIC)
        names_num = np.asarray(EP_NUMERIC, dtype=object)

        cat_names = []
        onehot = (
            pre.named_transformers_["cat"]
            .named_steps["onehot"]
        )
        cats = onehot.categories_
        for feature, levels in zip(EP_CATEGORICAL, cats):
            for level in levels:
                cat_names.append(f"{feature}_{level}")

        names = np.concatenate(
            [names_num, np.asarray(cat_names, dtype=object)]
        )

    support = selector.get_support()
    if len(names) != len(support):
        raise RuntimeError(
            f"Number of preprocessed features={len(names)},selector mask={len(support)}, mismatch"
        )

    rows = []
    selected_rank = 0
    for idx, (name, keep) in enumerate(zip(names, support)):
        name = str(name)

        # Group encoded features by their original predictor for SHAP aggregation.
        if name in EP_NUMERIC:
            parent = name
            kind = "numeric"
        elif name.startswith("Location_"):
            parent = "Location"
            kind = "onehot"
        elif name.startswith("Scene_"):
            parent = "Scene"
            kind = "onehot"
        else:
            # Account for alternative feature naming across scikit-learn versions.
            if "Location" in name:
                parent = "Location"
                kind = "onehot"
            elif "Scene" in name:
                parent = "Scene"
                kind = "onehot"
            else:
                parent = name
                kind = "unknown"

        if keep:
            selected_rank += 1

        rows.append({
            "Transformed_index": idx,
            "Transformed_feature": name,
            "Parent_feature": parent,
            "Feature_kind": kind,
            "Selected": bool(keep),
            "Selected_order": selected_rank if keep else np.nan,
        })

    return pd.DataFrame(rows)


# =============================================================================
# Checkpoint and resume support
# =============================================================================

def fold_key(season: str, fold: int) -> str:
    return f"{season}__F{int(fold)}"


def checkpoint_paths(run_dir: Path):
    checkpoint_dir = run_dir / "checkpoint"
    pred_dir = run_dir / "predictions"
    model_dir = run_dir / "models"
    feature_dir = run_dir / "feature_metadata"
    plot_ready_dir = run_dir / "plot_ready"

    for d in [checkpoint_dir, pred_dir, model_dir, feature_dir, plot_ready_dir]:
        d.mkdir(parents=True, exist_ok=True)

    return {
        "checkpoint_dir": checkpoint_dir,
        "checkpoint_csv": checkpoint_dir / "13-0_fold_results.csv",
        "pred_dir": pred_dir,
        "model_dir": model_dir,
        "feature_dir": feature_dir,
        "plot_ready_dir": plot_ready_dir,
    }


def fold_artifacts(paths: dict, season: str, fold: int):
    stem = f"13-0_{season}_F{fold}_XGB"
    return {
        "pred": paths["pred_dir"] / f"{stem}_predictions.csv",
        "model": paths["model_dir"] / f"{stem}.joblib",
        "features": paths["feature_dir"] / f"{stem}_features.csv",
    }


def load_checkpoint(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    df = pd.read_csv(path, encoding="utf-8-sig")
    return df


def checkpoint_fold_complete(
    ckpt: pd.DataFrame,
    paths: dict,
    season: str,
    fold: int,
) -> bool:
    if ckpt.empty:
        return False

    key = fold_key(season, fold)
    if "fold_key" not in ckpt.columns:
        return False

    if key not in set(ckpt["fold_key"].astype(str)):
        return False

    arts = fold_artifacts(paths, season, fold)
    return all(p.exists() for p in arts.values())


def append_checkpoint_row(
    checkpoint_csv: Path,
    row: dict,
):
    existing = load_checkpoint(checkpoint_csv)

    new = pd.DataFrame([row])

    if existing.empty:
        out = new
    else:
        existing = existing[
            existing["fold_key"].astype(str) != str(row["fold_key"])
        ]
        out = pd.concat([existing, new], ignore_index=True)

    out = out.sort_values(["Season", "Outer_fold"]).reset_index(drop=True)
    out.to_csv(
        checkpoint_csv,
        index=False,
        encoding="utf-8-sig",
    )


# =============================================================================
# Confirm input/configuration compatibility before resuming.
# =============================================================================

def build_run_config(input_file: Path, path12b: Path, smoke_test: bool):
    return {
        "script_version": SCRIPT_VERSION,
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "purpose": (
            "Final strict cross-fitted predictive evaluation of the selected common XGB family; "
            "NO RETUNING / NO RESELECTION"
        ),
        "input_file": str(input_file.resolve()),
        "input_sha256": sha256_file(input_file),
        "source_12B_excel": str(path12b.resolve()),
        "source_12B_sha256": sha256_file(path12b),
        "common_family": COMMON_FAMILY,
        "Stage 12B sheet 04_Best_Family_Configs, selected XGB configuration": "Stage 12B sheet 04_Best_Family_Configs, selected XGB configuration",
        "feature_set": "EP",
        "EP_features": EP_FEATURES,
        "target": TARGET,
        "raw_to_model": RAW_TO_MODEL,
        "random_state": RANDOM_STATE,
        "outer_fit_phase_offset": OUTER_FIT_PHASE_OFFSET,
        "model_n_jobs": MODEL_N_JOBS,
        "expected_n": EXPECTED_N,
        "expected_groups": EXPECTED_GROUPS,
        "smoke_test": bool(smoke_test),
    }


def validate_or_write_config(
    config_path: Path,
    current: dict,
):
    if config_path.exists():
        old = json.loads(config_path.read_text(encoding="utf-8"))

        keys = [
            "input_sha256",
            "source_12B_sha256",
            "common_family",
            "feature_set",
            "target",
            "smoke_test",
        ]

        mismatch = {
            k: (old.get(k), current.get(k))
            for k in keys
            if old.get(k) != current.get(k)
        }

        if mismatch:
            raise RuntimeError(
                "The current input or Stage 12B results differ from the existing checkpoint; "
                "resuming with mixed inputs is prohibited.\n"
                f"Mismatching items: {mismatch}\n"
                "Use --reset to start a fresh run after confirming the inputs."
            )
    else:
        config_path.write_text(
            json.dumps(
                current,
                ensure_ascii=False,
                indent=2,
            ),
            encoding="utf-8",
        )


# =============================================================================
# Evaluate one held-out day using frozen XGB configurations.
# =============================================================================

def run_one_fold(
    data: dict[str, pd.DataFrame],
    config_row: pd.Series,
    paths: dict,
):
    season = str(config_row["Season"])
    fold = int(config_row["Outer_fold"])
    held = pd.Timestamp(config_row["Held_out_date"]).normalize()

    season_df = data[season]

    dates = sorted(season_df["Date"].unique())
    expected_held = pd.Timestamp(dates[fold - 1]).normalize()

    # The fold index in Stage 12B follows chronological order; verify again.
    if held != expected_held:
        raise RuntimeError(
            f"{season} F{fold}: Stage 12B recorded held-out date={held.date()},"
            f"but sorted source dates imply F{fold}should be {expected_held.date()}."
        )

    train = (
        season_df[season_df["Date"] != held]
        .copy()
        .reset_index(drop=True)
    )
    test = (
        season_df[season_df["Date"] == held]
        .copy()
        .reset_index(drop=True)
    )

    if len(test) == 0:
        raise RuntimeError(
            f"{season} F{fold}: held-out day has no test observations"
        )

    train_groups = set(train["measurement_group"])
    test_groups = set(test["measurement_group"])
    overlap = train_groups & test_groups

    if overlap:
        raise RuntimeError(
            f"{season} F{fold}: outer train/test share "
            f"{len(overlap)} measurement groups"
        )

    raw_config = json.loads(config_row["Params_json"])
    selector, weight_mode, _ = decode_xgb_config(raw_config)
    seed = model_fit_seed(season, fold)

    print(
        f"\n[Evaluation] {season} F{fold} | held={held.date()} | "
        f"train={len(train)} | test={len(test)} | "
        f"selector={selector}% | weight={weight_mode} | seed={seed}"
    )
    print(
        f"       Stage 12B inner evidence: Robust_Min={float(config_row['Robust_Min']):.4f} | "
        f"Robust_Mean={float(config_row['Robust_Mean']):.4f} | "
        f"WorstDay={float(config_row['WorstDay_MacroF1']):.4f}"
    )

    t0 = time.time()

    pipe = make_xgb_pipeline(raw_config, seed)
    fit_pipeline(
        pipe,
        train[EP_FEATURES],
        train["y_model"].to_numpy(),
        weight_mode,
    )

    pred = pipe.predict(test[EP_FEATURES]).astype(int)
    proba = pipe.predict_proba(test[EP_FEATURES])

    ytrue = test["y_model"].to_numpy()
    metrics = metrics_bundle(ytrue, pred)

    if proba.shape != (len(test), 3):
        raise RuntimeError(
            f"{season} F{fold}: predict_proba shape={proba.shape}, expected (N, 3)"
        )

    duration = time.time() - t0

    # Construct record-level OOF predictions.
    pred_df = pd.DataFrame({
        "Season": season,
        "Outer_fold": fold,
        "Held_out_date": held.strftime("%Y-%m-%d"),
        "row_id": test["row_id"].astype(int).to_numpy(),
        "Date": test["Date"].dt.strftime("%Y-%m-%d").to_numpy(),
        "y_true_model": ytrue.astype(int),
        "y_pred_model": pred.astype(int),
        "y_true_raw": [MODEL_TO_RAW[int(v)] for v in ytrue],
        "y_pred_raw": [MODEL_TO_RAW[int(v)] for v in pred],
        "P_Dissatisfied": proba[:, 0],
        "P_Neutral": proba[:, 1],
        "P_Satisfied": proba[:, 2],
        "Model": COMMON_FAMILY,
        "Selector_percentile": selector,
        "Weight_mode": weight_mode,
        "Fit_seed": seed,
    }).sort_values("row_id")

    # Collect feature selection metadata.
    feature_meta = extract_selected_feature_metadata(pipe)
    feature_meta.insert(0, "Season", season)
    feature_meta.insert(1, "Outer_fold", fold)
    feature_meta.insert(2, "Held_out_date", held.strftime("%Y-%m-%d"))

    # Save fold-specific model and prediction artifacts.
    arts = fold_artifacts(paths, season, fold)
    pred_df.to_csv(arts["pred"], index=False, encoding="utf-8-sig")
    feature_meta.to_csv(
        arts["features"],
        index=False,
        encoding="utf-8-sig",
    )
    joblib.dump(pipe, arts["model"])

    row = {
        "fold_key": fold_key(season, fold),
        "Season": season,
        "Outer_fold": fold,
        "Held_out_date": held.strftime("%Y-%m-%d"),
        "Model": COMMON_FAMILY,
        "N_train": len(train),
        "N_test": len(test),
        "Train_groups": len(train_groups),
        "Test_groups": len(test_groups),
        "Group_overlap": len(overlap),
        "Selector_percentile": selector,
        "Weight_mode": weight_mode,
        "Fit_seed": seed,
        "Robust_Min": float(config_row["Robust_Min"]),
        "Robust_Mean": float(config_row["Robust_Mean"]),
        "WorstDay_MacroF1_inner": float(config_row["WorstDay_MacroF1"]),
        "Robust_WorstClassF1_inner": float(config_row["Robust_WorstClassF1"]),
        "Stability_MaxSD_inner": float(config_row["Stability_MaxSD"]),
        "Robust_BalancedAccuracy_inner": float(config_row["Robust_BalancedAccuracy"]),
        "Robust_Kappa_inner": float(config_row["Robust_Kappa"]),
        "Params_json": config_row["Params_json"],
        "Selected_transformed_features": int(feature_meta["Selected"].sum()),
        "Fit_seconds": float(duration),
        "Macro_F1": metrics["Macro_F1"],
        "Balanced_Accuracy": metrics["Balanced_Accuracy"],
        "Kappa": metrics["Kappa"],
        "Accuracy": metrics["Accuracy"],
        "Worst_Class_F1": metrics["Worst_Class_F1"],
        "Class1_Precision": metrics["Class1_Precision"],
        "Class1_Recall": metrics["Class1_Recall"],
        "Class1_F1": metrics["Class1_F1"],
        "Class1_Support": metrics["Class1_Support"],
        "Class2_Precision": metrics["Class2_Precision"],
        "Class2_Recall": metrics["Class2_Recall"],
        "Class2_F1": metrics["Class2_F1"],
        "Class2_Support": metrics["Class2_Support"],
        "Class3_Precision": metrics["Class3_Precision"],
        "Class3_Recall": metrics["Class3_Recall"],
        "Class3_F1": metrics["Class3_F1"],
        "Class3_Support": metrics["Class3_Support"],
        "Completed_at": datetime.now().isoformat(timespec="seconds"),
    }

    append_checkpoint_row(paths["checkpoint_csv"], row)

    print(
        f"       Outer XGB Macro-F1={metrics['Macro_F1']:.4f} | "
        f"BalAcc={metrics['Balanced_Accuracy']:.4f} | "
        f"Kappa={metrics['Kappa']:.4f} | "
        f"WorstClass={metrics['Worst_Class_F1']:.4f} | "
        f"{fmt_seconds(duration)}"
    )

    return row


# =============================================================================
# Aggregate cross-fitted results
# =============================================================================

def load_all_predictions(
    paths: dict,
    configs: pd.DataFrame,
) -> pd.DataFrame:
    dfs = []
    for _, r in configs.sort_values(["Season", "Outer_fold"]).iterrows():
        season = str(r["Season"])
        fold = int(r["Outer_fold"])
        p = fold_artifacts(paths, season, fold)["pred"]
        if not p.exists():
            raise FileNotFoundError(
                f"Missing fold prediction file: {p}"
            )
        dfs.append(pd.read_csv(p))

    out = pd.concat(dfs, ignore_index=True)
    out = out.sort_values(["Season", "row_id"]).reset_index(drop=True)
    return out


def build_oof_summary(
    pred_df: pd.DataFrame,
    fold_df: pd.DataFrame,
):
    summary_rows = []
    class_rows = []

    for season in ["Summer", "Winter"]:
        g = pred_df[pred_df["Season"] == season].copy()

        y = g["y_true_model"].astype(int).to_numpy()
        p = g["y_pred_model"].astype(int).to_numpy()
        m = metrics_bundle(y, p)

        ff = (
            fold_df[fold_df["Season"] == season]
            .copy()
            .sort_values("Outer_fold")
        )

        summary_rows.append({
            "Season": season,
            "N": len(g),
            "XGB_OOF_Macro_F1": m["Macro_F1"],
            "XGB_OOF_Balanced_Accuracy": m["Balanced_Accuracy"],
            "XGB_OOF_Kappa": m["Kappa"],
            "XGB_OOF_Accuracy": m["Accuracy"],
            "XGB_OOF_Worst_Class_F1": m["Worst_Class_F1"],
            "Dissatisfied_F1": m["Class1_F1"],
            "Neutral_F1": m["Class2_F1"],
            "Satisfied_F1": m["Class3_F1"],
            "Outer_fold_MacroF1_mean": float(ff["Macro_F1"].mean()),
            "Outer_fold_MacroF1_SD_descriptive": float(
                ff["Macro_F1"].std(ddof=1)
            ),
            "Outer_fold_MacroF1_worst": float(ff["Macro_F1"].min()),
            "Outer_fold_MacroF1_best": float(ff["Macro_F1"].max()),
            "Family": COMMON_FAMILY,
        })

        for cls in [1, 2, 3]:
            class_rows.append({
                "Season": season,
                "Class": cls,
                "Class_label": {
                    1: "Dissatisfied",
                    2: "Neutral",
                    3: "Satisfied",
                }[cls],
                "Precision": m[f"Class{cls}_Precision"],
                "Recall": m[f"Class{cls}_Recall"],
                "F1": m[f"Class{cls}_F1"],
                "Support": m[f"Class{cls}_Support"],
            })

    return pd.DataFrame(summary_rows), pd.DataFrame(class_rows)


def build_comparison_with_12b(
    path12b: Path,
    xgb_summary: pd.DataFrame,
):
    final12b = pd.read_excel(
        path12b,
        sheet_name="08_Pooled_OOF_Summary",
    )

    rows = []
    for season in ["Summer", "Winter"]:
        a = xgb_summary[xgb_summary["Season"] == season].iloc[0]
        b = final12b[final12b["Season"] == season].iloc[0]

        rows.append({
            "Season": season,
            "12B_LocalWinnerPipeline_Macro_F1": float(b["OOF_Macro_F1"]),
            "13-0_FinalXGB_Macro_F1": float(a["XGB_OOF_Macro_F1"]),
            "Delta_FinalXGB_minus_12BLocalWinner": (
                float(a["XGB_OOF_Macro_F1"]) - float(b["OOF_Macro_F1"])
            ),
            "12B_LocalWinnerPipeline_BalAcc": float(b["OOF_Balanced_Accuracy"]),
            "13-0_FinalXGB_BalAcc": float(a["XGB_OOF_Balanced_Accuracy"]),
            "12B_LocalWinnerPipeline_Kappa": float(b["OOF_Kappa"]),
            "13-0_FinalXGB_Kappa": float(a["XGB_OOF_Kappa"]),
            "12B_LocalWinnerPipeline_Accuracy": float(b["OOF_Accuracy"]),
            "13-0_FinalXGB_Accuracy": float(a["XGB_OOF_Accuracy"]),
            "Interpretation": (
                "12B local-winner pipeline is retained only as a model-selection benchmark; "
                "13-0 provides the formal strict cross-fitted performance of the selected common XGB family."
            ),
        })

    return pd.DataFrame(rows)


def reproducibility_check_for_xgb_winner_folds(
    path12b: Path,
    pred13: pd.DataFrame,
    fold13: pd.DataFrame,
):
    """Check reproducibility for folds where Stage 12B selected XGB locally.\n    On the same environment, the same parameters, seed, and preprocessing\n    should ordinarily give matching predictions. Other folds are excluded.\n    """
    outer12 = pd.read_excel(
        path12b,
        sheet_name="07_Outer_Fold_Performance",
    )
    pred12 = pd.read_excel(
        path12b,
        sheet_name="10_Outer_Predictions",
    )

    xgb_winners = outer12[
        outer12["Selected_model"].astype(str).str.strip() == "XGB"
    ].copy()

    rows = []

    for _, r in xgb_winners.iterrows():
        season = str(r["Season"])
        fold = int(r["Outer_fold"])

        a = pred12[
            (pred12["Season"] == season)
            & (pred12["Outer_fold"].astype(int) == fold)
        ][["row_id", "y_pred_model"]].copy()

        b = pred13[
            (pred13["Season"] == season)
            & (pred13["Outer_fold"].astype(int) == fold)
        ][["row_id", "y_pred_model"]].copy()

        merged = a.merge(
            b,
            on="row_id",
            suffixes=("_12B", "_13_0"),
            how="outer",
            indicator=True,
        )

        same_rows = bool((merged["_merge"] == "both").all())
        if same_rows:
            pred_match = float(
                (
                    merged["y_pred_model_12B"].astype(int)
                    == merged["y_pred_model_13_0"].astype(int)
                ).mean()
            )
        else:
            pred_match = np.nan

        f13 = fold13[
            (fold13["Season"] == season)
            & (fold13["Outer_fold"].astype(int) == fold)
        ].iloc[0]

        macro_diff = float(f13["Macro_F1"]) - float(r["Macro_F1"])

        rows.append({
            "Season": season,
            "Outer_fold": fold,
            "Held_out_date": pd.Timestamp(r["Held_out_date"]).strftime("%Y-%m-%d"),
            "12B_Selected_model": "XGB",
            "Prediction_row_set_identical": same_rows,
            "Prediction_exact_match_fraction": pred_match,
            "12B_Macro_F1": float(r["Macro_F1"]),
            "13-0_Macro_F1": float(f13["Macro_F1"]),
            "Macro_F1_difference": macro_diff,
            "Note": (
                "Expected to reproduce because 12B winner was already XGB; "
                "small discrepancies can occur after library/version changes."
            ),
        })

    return pd.DataFrame(rows)


def combine_feature_metadata(
    paths: dict,
    configs: pd.DataFrame,
):
    dfs = []
    for _, r in configs.iterrows():
        p = fold_artifacts(
            paths,
            str(r["Season"]),
            int(r["Outer_fold"]),
        )["features"]
        dfs.append(pd.read_csv(p))
    return pd.concat(dfs, ignore_index=True)


# =============================================================================
# Export plot-ready data for the manuscript figures and tables.
# =============================================================================

def export_plot_ready_v2(
    paths: dict,
    fold_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    class_df: pd.DataFrame,
    pred_df: pd.DataFrame,
):
    """Export data used by manuscript Fig. 4, Fig. 5, and Table 4."""
    out_dir = paths["plot_ready_dir"]

    # Fig. 4: held-out-day Macro-F1 and pooled OOF Macro-F1.
    pooled_map = summary_df.set_index("Season")["XGB_OOF_Macro_F1"].to_dict()
    fig4 = fold_df[[
        "Season", "Outer_fold", "Held_out_date", "N_train", "N_test",
        "Macro_F1", "Balanced_Accuracy", "Kappa", "Accuracy", "Worst_Class_F1"
    ]].copy()
    fig4["Pooled_OOF_Macro_F1"] = fig4["Season"].map(pooled_map)
    fig4 = fig4.sort_values(["Season", "Outer_fold"]).reset_index(drop=True)
    fig4.to_csv(out_dir / "Fig4_DayToDay_Generalization_Ready.csv", index=False, encoding="utf-8-sig")

    # Fig. 5: class-specific predictive performance.
    fig5_class = class_df.copy().sort_values(["Season", "Class"]).reset_index(drop=True)
    fig5_class.to_csv(out_dir / "Fig5_ClassPerformance_Ready.csv", index=False, encoding="utf-8-sig")

    # Fig.5:row-normalized confusion matrix
    cm_rows = []
    labels = [0, 1, 2]
    label_names = {0: "Dissatisfied", 1: "Neutral", 2: "Satisfied"}
    for season in ["Summer", "Winter"]:
        g = pred_df[pred_df["Season"] == season].copy()
        y_true = g["y_true_model"].astype(int).to_numpy()
        y_pred = g["y_pred_model"].astype(int).to_numpy()
        cm_count = confusion_matrix(y_true, y_pred, labels=labels)
        row_sum = cm_count.sum(axis=1, keepdims=True)
        cm_norm = np.divide(cm_count, row_sum, out=np.zeros_like(cm_count, dtype=float), where=row_sum != 0)
        for i, true_cls in enumerate(labels):
            for j, pred_cls in enumerate(labels):
                cm_rows.append({
                    "Season": season,
                    "True_class_model": true_cls,
                    "True_class": label_names[true_cls],
                    "Predicted_class_model": pred_cls,
                    "Predicted_class": label_names[pred_cls],
                    "Count": int(cm_count[i, j]),
                    "Row_normalized": float(cm_norm[i, j]),
                    "Row_normalized_pct": float(cm_norm[i, j] * 100.0),
                })
    fig5_cm = pd.DataFrame(cm_rows)
    fig5_cm.to_csv(out_dir / "Fig5_ConfusionMatrix_Ready.csv", index=False, encoding="utf-8-sig")

    # Table 4: pooled predictive performance.
    table4 = summary_df.copy()
    table4.to_csv(out_dir / "Table4_OverallPerformance_Ready.csv", index=False, encoding="utf-8-sig")

    # Export supplementary descriptive fold metrics.
    fold_df.sort_values(["Season", "Outer_fold"]).to_csv(
        out_dir / "HeldOut_Fold_Performance_Ready.csv", index=False, encoding="utf-8-sig"
    )

    return fig4, fig5_class, fig5_cm, table4


# =============================================================================
# Excel output
# =============================================================================

def style_excel(path: Path):
    wb = load_workbook(path)

    fill = PatternFill("solid", fgColor="D4ECEF")
    header_font = Font(
        name="Times New Roman",
        size=10,
        bold=True,
    )
    body_font = Font(
        name="Times New Roman",
        size=10,
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        ws.auto_filter.ref = ws.dimensions

        for cell in ws[1]:
            cell.fill = fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.font = body_font
                c.alignment = Alignment(
                    horizontal="center",
                    vertical="center",
                    wrap_text=True,
                )
                if isinstance(c.value, float):
                    c.number_format = "0.000"

        for idx, col_cells in enumerate(
            ws.columns, start=1
        ):
            max_len = 0
            for c in col_cells:
                if c.value is not None:
                    max_len = max(max_len, len(str(c.value)))
            ws.column_dimensions[
                get_column_letter(idx)
            ].width = min(max(max_len + 2, 10), 32)

    wb.save(path)


def export_excel(
    run_dir: Path,
    run_config: dict,
    xgb_configs: pd.DataFrame,
    fold_df: pd.DataFrame,
    summary_df: pd.DataFrame,
    class_df: pd.DataFrame,
    pred_df: pd.DataFrame,
    comparison_df: pd.DataFrame,
    repro_df: pd.DataFrame,
    feature_meta: pd.DataFrame,
):
    output = run_dir / OUTPUT_XLSX_NAME

    config_table = pd.DataFrame(
        {
            "Item": [
                "Step",
                "Purpose",
                "Retuning",
                "Common family",
                "Input",
                "12B source",
                "Feature set",
                "Target",
                "Outer validation",
                "Parameter source",
                "Random state",
                "N Summer",
                "N Winter",
                "Groups Summer",
                "Groups Winter",
            ],
            "Value": [
                "13-0 XGB-only cross-fitted qualification",
                "Verify strict cross-day performance of the common interpretation family before SHAP",
                "NO",
                "XGB",
                run_config["input_file"],
                run_config["source_12B_excel"],
                "EP",
                TARGET,
                "4-fold leave-one-day-out separately by season",
                "Stage 12B sheet 04_Best_Family_Configs, selected XGB configuration",
                RANDOM_STATE,
                EXPECTED_N["Summer"],
                EXPECTED_N["Winter"],
                EXPECTED_GROUPS["Summer"],
                EXPECTED_GROUPS["Winter"],
            ],
        }
    )

    param_cols = [
        "Season",
        "Outer_fold",
        "Held_out_date",
        "Model",
        "StageA_rank",
        "StageA_trial_number",
        "StageA_inner_Macro_F1",
        "FreshSGKF_MacroF1_mean",
        "FreshSGKF_MacroF1_SD",
        "InnerDayLODO_OOF_MacroF1",
        "WorstDay_MacroF1",
        "Day_MacroF1_SD",
        "Robust_Min",
        "Robust_Mean",
        "Robust_WorstClassF1",
        "Stability_MaxSD",
        "Robust_BalancedAccuracy",
        "Robust_Kappa",
        "Params_json",
    ]
    param_cols = [
        c for c in param_cols if c in xgb_configs.columns
    ]
    params_out = xgb_configs[param_cols].copy()

    # Use a consistent date representation.
    if "Held_out_date" in params_out.columns:
        params_out["Held_out_date"] = pd.to_datetime(
            params_out["Held_out_date"]
        ).dt.strftime("%Y-%m-%d")

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:
        config_table.to_excel(
            writer,
            sheet_name="00_Run_Configuration",
            index=False,
        )
        params_out.to_excel(
            writer,
            sheet_name="01_XGB_Fold_Configurations",
            index=False,
        )
        fold_df.to_excel(
            writer,
            sheet_name="02_XGB_Outer_Performance",
            index=False,
        )
        summary_df.to_excel(
            writer,
            sheet_name="03_XGB_Pooled_OOF",
            index=False,
        )
        class_df.to_excel(
            writer,
            sheet_name="04_XGB_Classwise_OOF",
            index=False,
        )
        pred_df.to_excel(
            writer,
            sheet_name="05_XGB_OOF_Predictions",
            index=False,
        )
        comparison_df.to_excel(
            writer,
            sheet_name="06_StageB_Local_Winners",
            index=False,
        )
        repro_df.to_excel(
            writer,
            sheet_name="07_XGB_Reproducibility_Check",
            index=False,
        )
        feature_meta.to_excel(
            writer,
            sheet_name="08_Selected_Feature_Metadata",
            index=False,
        )

    style_excel(output)
    return output


# =============================================================================
# Resolve input and output paths
# =============================================================================

def resolve_paths(args):
    base_dir = (
        Path(args.base_dir).expanduser().resolve()
        if args.base_dir
        else DEFAULT_BASE_DIR
    )

    input_file = (
        Path(args.input).expanduser().resolve()
        if args.input
        else base_dir / INPUT_NAME
    )

    dir12b = base_dir / DIR_12B_NAME
    path12b = (
        Path(args.result12b).expanduser().resolve()
        if args.result12b
        else resolve_xlsx(
            dir12b / FILE_12B_NAME,
            dir12b,
            "12B",
        )
    )

    suffix = "_SMOKE_TEST" if args.smoke_test else ""
    run_dir = base_dir / f"{RUN_DIR_NAME}{suffix}"

    if args.reset and run_dir.exists():
        print(f"[RESET] Removing previous output directory: {run_dir}")
        shutil.rmtree(run_dir)

    run_dir.mkdir(parents=True, exist_ok=True)

    if not input_file.exists():
        raise FileNotFoundError(
            f"Full analytical workbook not found:\n{input_file}"
        )
    if not path12b.exists():
        raise FileNotFoundError(
            f"Stage 12B results workbook not found:\n{path12b}"
        )

    return base_dir, input_file, path12b, run_dir


# =============================================================================
# Main workflow
# =============================================================================

def main():
    parser = argparse.ArgumentParser(
        description=(
            "Final XGB model: use frozen Stage 12B inner-selected parameters to "
            "evaluate held-out days without further tuning."
        )
    )
    parser.add_argument(
        "--base-dir",
        default=None,
        help="Base directory (defaults to the script directory)",
    )
    parser.add_argument(
        "--input",
        default=None,
        help="Path to the full analytical Excel workbook",
    )
    parser.add_argument(
        "--result12b",
        default=None,
        help="Path to Stage 12B result workbook",
    )
    parser.add_argument(
        "--smoke-test",
        action="store_true",
        help="Evaluate only summer fold 1 in an isolated smoke-test directory.",
    )
    parser.add_argument(
        "--reset",
        action="store_true",
        help="Remove the final-XGB output directory and restart.",
    )

    args = parser.parse_args()

    base_dir, input_file, path12b, run_dir = resolve_paths(args)

    # Start run logging.
    log_path = run_dir / LOG_NAME
    log_file = log_path.open("a", encoding="utf-8", buffering=1)
    original_stdout = sys.stdout
    original_stderr = sys.stderr
    sys.stdout = Tee(original_stdout, log_file)
    sys.stderr = Tee(original_stderr, log_file)

    try:
        print("\n" + "=" * 108)
        print("Final XGB evaluation using frozen configurations / cross-fitted LODO")
        print("=" * 108)
        print(f"Started: {datetime.now()}")
        print(f"Base dir:{base_dir}")
        print(f"Input dataset: {input_file}")
        print(f"Stage 12B results: {path12b}")
        print(f"Output directory: {run_dir}")
        print(f"Smoke test:{args.smoke_test}")
        print("Retuning: False")
        print("Configuration source: Stage 12B inner/stability-selected XGB parameters")
        print("Final evaluation is separate from Stage 12B local-winner screening.")
        print("=" * 108)

        # Avoid mixing inputs/configurations from different runs.
        run_config = build_run_config(
            input_file,
            path12b,
            args.smoke_test,
        )
        config_path = run_dir / CONFIG_JSON_NAME
        validate_or_write_config(
            config_path,
            run_config,
        )

        # Read data and frozen configurations.
        data = load_data(input_file)
        xgb_configs, common_table = load_12b_xgb_configs(path12b)

        # Check held-out dates against the source dataset.
        for season in ["Summer", "Winter"]:
            dates_data = [
                pd.Timestamp(x).normalize()
                for x in sorted(data[season]["Date"].unique())
            ]
            dates_cfg = [
                pd.Timestamp(x).normalize()
                for x in (
                    xgb_configs[xgb_configs["Season"] == season]
                    .sort_values("Outer_fold")["Held_out_date"]
                    .tolist()
                )
            ]
            if dates_data != dates_cfg:
                raise RuntimeError(
                    f"{season}: Stage 12B XGB configuration dates differ from the four source survey dates.\n"
                    f"Data={dates_data}\n12B={dates_cfg}"
                )

        if args.smoke_test:
            configs_run = xgb_configs[
                (xgb_configs["Season"] == "Summer")
                & (xgb_configs["Outer_fold"].astype(int) == 1)
            ].copy()
        else:
            configs_run = xgb_configs.copy()

        paths = checkpoint_paths(run_dir)
        ckpt = load_checkpoint(paths["checkpoint_csv"])

        completed = 0
        for _, r in configs_run.iterrows():
            if checkpoint_fold_complete(
                ckpt,
                paths,
                str(r["Season"]),
                int(r["Outer_fold"]),
            ):
                completed += 1

        total = len(configs_run)

        print(
            f"\n[Run plan] folds={total} | "
            f"Previously completed={completed} | "
            f"Pending={total-completed}"
        )
        print_progress(completed, total, time.time())

        start = time.time()

        # Evaluate all requested outer folds.
        for _, r in configs_run.sort_values(
            ["Season", "Outer_fold"]
        ).iterrows():
            season = str(r["Season"])
            fold = int(r["Outer_fold"])

            ckpt = load_checkpoint(paths["checkpoint_csv"])

            if checkpoint_fold_complete(
                ckpt,
                paths,
                season,
                fold,
            ):
                print(
                    f"\n[RESUME: skipped] {season} F{fold} already completed"
                )
                continue

            run_one_fold(
                data=data,
                config_row=r,
                paths=paths,
            )

            ckpt = load_checkpoint(paths["checkpoint_csv"])
            done_now = sum(
                checkpoint_fold_complete(
                    ckpt,
                    paths,
                    str(rr["Season"]),
                    int(rr["Outer_fold"]),
                )
                for _, rr in configs_run.iterrows()
            )
            print_progress(done_now, total, start)

        # Smoke tests do not generate formal pooled results.
        if args.smoke_test:
            print("\n" + "=" * 108)
            print("Final XGB smoke test complete; not intended for publication.")
            print("Rerun without --smoke-test for the formal analysis.")
            print(f"Models: {paths['model_dir']}")
            print(f"Predictions: {paths['pred_dir']}")
            print(f"Finished: {datetime.now()}")
            print("=" * 108)
            return

        # Formal evaluation requires all eight folds.
        ckpt = load_checkpoint(paths["checkpoint_csv"])
        done = sum(
            checkpoint_fold_complete(
                ckpt,
                paths,
                str(r["Season"]),
                int(r["Outer_fold"]),
            )
            for _, r in xgb_configs.iterrows()
        )
        if done != 8:
            raise RuntimeError(
                f"Formal pooled evaluation requires all eight folds; completed={done}/8"
            )

        fold_df = (
            ckpt.sort_values(["Season", "Outer_fold"])
            .reset_index(drop=True)
        )

        pred_df = load_all_predictions(
            paths,
            xgb_configs,
        )

        # Verify that all 1,001 OOF predictions are present.
        if len(pred_df) != 1001:
            raise RuntimeError(
                f"Expected 1,001 XGB OOF predictions; found={len(pred_df)}"
            )
        if int((pred_df["Season"] == "Summer").sum()) != 477:
            raise RuntimeError("Summer XGB OOF predictions do not contain all 477 observations.")
        if int((pred_df["Season"] == "Winter").sum()) != 524:
            raise RuntimeError("Winter XGB OOF predictions do not contain all 524 observations.")

        # Each original row identifier must appear exactly once per season.
        for season, n in EXPECTED_N.items():
            g = pred_df[pred_df["Season"] == season]
            if g["row_id"].nunique() != n:
                raise RuntimeError(
                    f"{season}: OOF row identifiers do not occur exactly once per observation"
                )

        summary_df, class_df = build_oof_summary(
            pred_df,
            fold_df,
        )

        comparison_df = build_comparison_with_12b(
            path12b,
            summary_df,
        )

        repro_df = reproducibility_check_for_xgb_winner_folds(
            path12b,
            pred_df,
            fold_df,
        )

        feature_meta = combine_feature_metadata(
            paths,
            xgb_configs,
        )

        # Export data for current manuscript figures and tables.
        export_plot_ready_v2(
            paths=paths,
            fold_df=fold_df,
            summary_df=summary_df,
            class_df=class_df,
            pred_df=pred_df,
        )

        output_excel = export_excel(
            run_dir=run_dir,
            run_config=run_config,
            xgb_configs=xgb_configs,
            fold_df=fold_df,
            summary_df=summary_df,
            class_df=class_df,
            pred_df=pred_df,
            comparison_df=comparison_df,
            repro_df=repro_df,
            feature_meta=feature_meta,
        )

        # Export master tables in CSV format for downstream figures and reporting.
        summary_df.to_csv(
            run_dir / "Table4_OverallPerformance_Ready.csv",
            index=False,
            encoding="utf-8-sig",
        )
        class_df.to_csv(
            run_dir / "Fig5_ClassPerformance_Ready.csv",
            index=False,
            encoding="utf-8-sig",
        )

        print("\n" + "=" * 108)
        print("Final XGB LODO evaluation completed.")
        print("=" * 108)

        for _, r in summary_df.iterrows():
            print(
                f"{r['Season']}: "
                f"Macro-F1={r['XGB_OOF_Macro_F1']:.4f} | "
                f"BalAcc={r['XGB_OOF_Balanced_Accuracy']:.4f} | "
                f"Kappa={r['XGB_OOF_Kappa']:.4f} | "
                f"Accuracy={r['XGB_OOF_Accuracy']:.4f} | "
                f"Class F1="
                f"{r['Dissatisfied_F1']:.4f}/"
                f"{r['Neutral_F1']:.4f}/"
                f"{r['Satisfied_F1']:.4f} | "
                f"Fold SD={r['Outer_fold_MacroF1_SD_descriptive']:.4f}"
            )

        print("\nComparison against Stage 12B selected local-winner pipelines:")
        for _, r in comparison_df.iterrows():
            print(
                f"{r['Season']}: 12B={r['12B_LocalWinnerPipeline_Macro_F1']:.4f} "
                f"| XGB-only={r['13-0_FinalXGB_Macro_F1']:.4f} "
                f"| Δ={r['Delta_FinalXGB_minus_12BLocalWinner']:+.4f}"
            )

        if len(repro_df):
            print("\nReproducibility checks for folds where Stage 12B selected XGB:")
            for _, r in repro_df.iterrows():
                print(
                    f"{r['Season']} F{int(r['Outer_fold'])}: "
                    f"prediction exact match="
                    f"{r['Prediction_exact_match_fraction']:.3f} | "
                    f"Macro-F1 diff={r['Macro_F1_difference']:+.8f}"
                )

        print(f"\nResults workbook: {output_excel}")
        print(f"Eight XGB pipelines: {paths['model_dir']}")
        print(f"OOF predictions: {paths['pred_dir']}")
        print(f"Feature metadata: {paths['feature_dir']}")
        print(
            "\nNext, calculate held-out-day cross-fitted SHAP using these eight pipelines; "
            "each pipeline explains only its own held-out-day samples; "
            "no additional hyperparameter tuning."
        )
        print(f"Finished: {datetime.now()}")
        print("=" * 108)

    finally:
        sys.stdout = original_stdout
        sys.stderr = original_stderr
        log_file.close()


if __name__ == "__main__":
    main()
